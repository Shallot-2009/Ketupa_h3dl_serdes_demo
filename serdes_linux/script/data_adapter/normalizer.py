"""Convert an analyzed source table to Ketupa's stable workbook contracts."""

from __future__ import annotations

import hashlib
import json
import os
import re
import tempfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from .classifier import analyze
from .models import AdapterError, Analysis, NoNormalizedRows
from .network_rules import map_network
from .tabular import cell


NETLIST_HEADERS = ("Net Name", "PIN 1", "PIN 2", "Signal Type", "M Group", "Chip Partition")
PLACEMENT_HEADERS = ("RefDes", "Role", "Side", "Mirror", "X", "Y", "Rotation", "Package")
IR_HEADERS = (
    "Physical Net",
    "Logical ID",
    "Family",
    "Topology",
    "Group",
    "Signal Type",
    "Polarity",
    "Segment",
    "Rule ID",
    "Source Sheet",
    "Source Row",
    "Endpoint A RefDes",
    "Endpoint B RefDes",
)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while chunk := stream.read(4 * 1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def _split_pins(value: str, profile: dict) -> list[str]:
    pattern = profile.get("pin", {}).get("separator_regex", r"[\s,;，；|]+")
    return [token for token in re.split(pattern, value.strip()) if token]


def _pin_refdes(pin: str, profile: dict, row_number: int) -> str:
    pattern = profile.get("pin", {}).get("token_regex", r"^(?P<refdes>[^.\s]+)\.(?P<pin>\S+)$")
    match = re.fullmatch(pattern, pin)
    if not match or "refdes" not in match.groupdict():
        raise AdapterError(f"Row {row_number}: invalid RefDes.Pin token {pin!r}")
    return match.group("refdes")


def _partition_all_pins(pins: list[str], profile: dict, row_number: int) -> tuple[list[str], list[str]]:
    rules = profile.get("endpoint_rules", [])
    if rules:
        output = {"a": [], "b": []}
        unresolved = False
        for pin in pins:
            refdes = _pin_refdes(pin, profile, row_number)
            targets = {
                str(rule["target"]).casefold()
                for rule in rules
                if re.search(rule["refdes_pattern"], refdes, re.IGNORECASE)
            }
            if len(targets) != 1 or not targets <= {"a", "b"}:
                unresolved = True
                break
            output[targets.pop()].append(pin)
        if not unresolved and output["a"] and output["b"]:
            return output["a"], output["b"]

    order_rules = profile.get("endpoint_order_rules", [])
    by_component: dict[str, list[str]] = {}
    for pin in pins:
        refdes = _pin_refdes(pin, profile, row_number)
        by_component.setdefault(refdes, []).append(pin)
    ranked: list[tuple[int, str, list[str]]] = []
    for refdes, component_pins in by_component.items():
        ranks = {
            int(rule["rank"])
            for rule in order_rules
            if re.search(rule["refdes_pattern"], refdes, re.IGNORECASE)
        }
        if len(ranks) > 1:
            raise AdapterError(
                f"Row {row_number}: endpoint rank for RefDes {refdes!r} is ambiguous; "
                "add an endpoint rule to the selected profile"
            )
        ranked.append((ranks.pop() if ranks else 1000, refdes.casefold(), component_pins))
    ranked.sort(key=lambda item: (item[0], item[1]))
    if len(ranked) < 2 or ranked[0][0] == ranked[1][0]:
        raise AdapterError(
            f"Row {row_number}: endpoints cannot be ordered safely; add project endpoint rules"
        )
    endpoint_a = list(ranked[0][2])
    endpoint_b = [pin for _, _, values in ranked[1:] for pin in values]
    if not endpoint_a or not endpoint_b:
        raise AdapterError(f"Row {row_number}: both canonical endpoints are required")
    return endpoint_a, endpoint_b


def _mapped(analysis: Analysis, row: tuple[Any, ...], field: str) -> str:
    return cell(row, analysis.table.mapping.get(field))


def _normalize_netlist(
    analysis: Analysis,
    profile: dict,
    family: str,
    topology: str,
    groups: str = "auto",
) -> tuple[list[tuple[str, ...]], list[tuple[str, ...]], dict[str, Any]]:
    rows: list[tuple[str, ...]] = []
    ir_rows: list[tuple[str, ...]] = []
    seen: set[str] = set()
    group_polarities: dict[str, Counter[str]] = defaultdict(Counter)
    rules_used: Counter[str] = Counter()
    skipped_blank = 0
    skipped_unmatched = 0
    contract = profile.get("network_contracts", {}).get(family, {}).get(topology, {})
    industry_contract = isinstance(contract, dict) and contract.get("recognizer") == "industry"
    source_names = {
        _mapped(analysis, row, "net_name").strip()
        for row in analysis.table.rows
        if _mapped(analysis, row, "net_name").strip()
    }
    normalized_names = {value.casefold() for value in source_names}
    starred_clock_bases = {
        value[:-1].strip().casefold()
        for value in source_names
        if family == "pll_clk"
        and value.endswith("*")
        and value[:-1].strip().casefold() in normalized_names
    }
    tc_pair_bases = {
        value[:-2].casefold()
        for value in source_names
        if value.upper().endswith("_T")
        and (value[:-2] + "_C").casefold() in normalized_names
    }
    for offset, raw in enumerate(analysis.table.rows, start=analysis.table.header_row + 1):
        name = _mapped(analysis, raw, "net_name")
        if not name:
            skipped_blank += 1
            continue
        key = name.casefold()
        if key in seen:
            raise AdapterError(f"Row {offset}: duplicate network {name!r}")
        seen.add(key)
        classification_name = name
        if family == "pll_clk" and name.casefold() in starred_clock_bases:
            classification_name = name + "_P"
        elif (
            family == "pll_clk"
            and name.endswith("*")
            and name[:-1].strip().casefold() not in starred_clock_bases
        ):
            classification_name = name[:-1]
        elif name.upper().endswith("_T") and name[:-2].casefold() in tc_pair_bases:
            classification_name = name[:-2] + "_P"
        elif name.upper().endswith("_C") and name[:-2].casefold() in tc_pair_bases:
            classification_name = name[:-2] + "_N"
        if industry_contract:
            from .industry_patterns import identify

            preliminary = identify(
                name=classification_name,
                family=family,
                topology=topology,
                endpoint_a=["AUTO.0"],
                source_signal_type=_mapped(analysis, raw, "signal_type"),
                source_group=_mapped(analysis, raw, "group"),
                source_partition=(
                    _mapped(analysis, raw, "partition")
                    or profile.get("defaults", {}).get("partition", "")
                ),
            )
            if preliminary is None and str(contract.get("unmatched_policy", "reject")).casefold() == "skip":
                skipped_unmatched += 1
                continue
        endpoint_a = _split_pins(_mapped(analysis, raw, "endpoint_a"), profile)
        endpoint_b = _split_pins(_mapped(analysis, raw, "endpoint_b"), profile)
        if not endpoint_a and not endpoint_b:
            all_pins = _split_pins(_mapped(analysis, raw, "all_pins"), profile)
            if industry_contract:
                components = {
                    _pin_refdes(pin, profile, offset) for pin in all_pins
                }
                if len(components) < 2:
                    skipped_unmatched += 1
                    continue
            endpoint_a, endpoint_b = _partition_all_pins(all_pins, profile, offset)
        if not endpoint_a or not endpoint_b:
            raise AdapterError(f"Row {offset}: both PIN 1 and PIN 2 are required")
        for pin in (*endpoint_a, *endpoint_b):
            _pin_refdes(pin, profile, offset)
        identity = map_network(
            name=classification_name,
            endpoint_a=endpoint_a,
            endpoint_b=endpoint_b,
            source_signal_type=_mapped(analysis, raw, "signal_type"),
            source_group=_mapped(analysis, raw, "group"),
            source_partition=(
                _mapped(analysis, raw, "partition")
                or profile.get("defaults", {}).get("partition", "")
            ),
            family=family,
            topology=topology,
            profile=profile,
        )
        if identity is None:
            skipped_unmatched += 1
            continue
        if not _group_selected(identity.group, groups):
            continue
        signal_type = identity.signal_type
        group = identity.group
        partition = identity.partition
        rules_used[identity.rule_id] += 1
        if signal_type.upper() in {"DIFF_P", "DIFF_N"}:
            group_polarities[group][signal_type.upper()[-1]] += 1
        rows.append((name, " ".join(endpoint_a), " ".join(endpoint_b), signal_type, group, partition))
        ir_rows.append(
            (
                name,
                identity.logical_id,
                family,
                topology,
                group,
                signal_type,
                identity.polarity,
                identity.segment,
                identity.rule_id,
                analysis.table.sheet,
                str(offset),
                ",".join(dict.fromkeys(_pin_refdes(pin, profile, offset) for pin in endpoint_a)),
                ",".join(dict.fromkeys(_pin_refdes(pin, profile, offset) for pin in endpoint_b)),
            )
        )
    if not rows:
        raise NoNormalizedRows("No data rows were normalized")
    incomplete = {
        group: dict(counts)
        for group, counts in group_polarities.items()
        if set(counts) != {"P", "N"} or counts["P"] != counts["N"]
    }
    if incomplete:
        raise AdapterError("Differential groups are incomplete or unbalanced: " + json.dumps(incomplete, ensure_ascii=False))
    return rows, ir_rows, {
        "rows": len(rows),
        "skipped_blank_rows": skipped_blank,
        "skipped_unmatched_rows": skipped_unmatched,
        "differential_groups": len(group_polarities),
        "network_rules_used": dict(sorted(rules_used.items())),
        "logical_identities": len({row[1] for row in ir_rows}),
    }


def _group_selected(group: str, selection: str) -> bool:
    text = str(selection or "auto").strip()
    if text.casefold() in {"", "auto", "all"}:
        return True
    normalized = str(group).strip().casefold()
    group_number_match = re.search(r"([0-9]+)$", normalized)
    group_number = int(group_number_match.group(1)) if group_number_match else None
    for raw in text.split(","):
        token = raw.strip().casefold()
        if not token:
            continue
        if token == normalized or token in normalized:
            return True
        compact = re.sub(r"^(?:m|lane|rx|tx)", "", token)
        if "-" in compact and group_number is not None:
            start, end = compact.split("-", 1)
            if start.isdigit() and end.isdigit() and int(start) <= group_number <= int(end):
                return True
        elif compact.isdigit() and group_number == int(compact):
            return True
    return False


def _role_for(refdes: str, profile: dict, row_number: int) -> str:
    matches = {
        str(rule["role"])
        for rule in profile.get("component_role_rules", [])
        if re.search(rule["refdes_pattern"], refdes, re.IGNORECASE)
    }
    if len(matches) > 1:
        raise AdapterError(f"Row {row_number}: conflicting component roles for {refdes}")
    return next(iter(matches), "")


def _normalize_placement(analysis: Analysis, profile: dict) -> tuple[list[tuple[str, ...]], dict[str, Any]]:
    rows: list[tuple[str, ...]] = []
    seen: set[str] = set()
    for offset, raw in enumerate(analysis.table.rows, start=analysis.table.header_row + 1):
        refdes = _mapped(analysis, raw, "refdes")
        if not refdes:
            continue
        key = refdes.casefold()
        if key in seen:
            raise AdapterError(f"Row {offset}: duplicate RefDes {refdes!r}")
        seen.add(key)
        role = _mapped(analysis, raw, "role") or _role_for(refdes, profile, offset)
        side = _mapped(analysis, raw, "side").upper()
        mirror = _mapped(analysis, raw, "mirror").upper()
        x = _mapped(analysis, raw, "x")
        y = _mapped(analysis, raw, "y")
        rotation = _mapped(analysis, raw, "rotation")
        package = _mapped(analysis, raw, "package")
        required = {"Role": role, "Side": side, "X": x, "Y": y, "Rotation": rotation, "Package": package}
        missing = [name for name, value in required.items() if not value]
        if missing:
            raise AdapterError(f"Row {offset}: missing placement values: {', '.join(missing)}")
        if side not in {"TOP", "BOTTOM"}:
            raise AdapterError(f"Row {offset}: Side must be TOP or BOTTOM, found {side!r}")
        if not mirror:
            mirror = "YES" if side == "BOTTOM" else "NO"
        try:
            float(x); float(y); float(rotation)
        except ValueError as exc:
            raise AdapterError(f"Row {offset}: X/Y/Rotation must be numeric") from exc
        rows.append((refdes, role, side, mirror, x, y, rotation, package))
    if not rows:
        raise AdapterError("No placement rows were normalized")
    return rows, {"rows": len(rows), "roles": dict(Counter(row[1] for row in rows))}


def _write_workbook(
    output: Path,
    headers: tuple[str, ...],
    rows: list[tuple[str, ...]],
    ir_rows: list[tuple[str, ...]] | None = None,
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ALL" if headers == NETLIST_HEADERS else "Placement"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    fill = PatternFill("solid", fgColor="17365D")
    for cell_value in sheet[1]:
        cell_value.fill = fill
        cell_value.font = Font(color="FFFFFF", bold=True)
        cell_value.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(70, max(12, max(len(str(item.value or "")) for item in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width
    if ir_rows:
        ir_sheet = workbook.create_sheet("_Ketupa_IR")
        ir_sheet.append(IR_HEADERS)
        for row in ir_rows:
            ir_sheet.append(row)
        for header in ir_sheet[1]:
            header.fill = fill
            header.font = Font(color="FFFFFF", bold=True)
        ir_sheet.freeze_panes = "A2"
        ir_sheet.auto_filter.ref = ir_sheet.dimensions
        ir_sheet.sheet_state = "hidden"
    output.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary = tempfile.mkstemp(prefix=output.stem + ".", suffix=".tmp.xlsx", dir=output.parent)
    os.close(handle)
    try:
        workbook.save(temporary)
        os.replace(temporary, output)
    finally:
        workbook.close()
        if os.path.exists(temporary):
            os.unlink(temporary)


def normalize(
    source: str | Path,
    output: str | Path,
    profile: dict,
    profile_sha256: str,
    family: str = "auto",
    topology: str = "auto",
    kind: str = "auto",
    sheet: str = "",
    force: bool = False,
    groups: str = "auto",
) -> dict[str, Any]:
    analysis = analyze(source, profile, profile_sha256, sheet=sheet)
    selected_kind = analysis.kind.name if kind == "auto" else kind
    selected_family = family
    selected_topology = topology
    if selected_kind not in {"netlist", "placement"}:
        raise AdapterError("Kind is ambiguous; pass --kind netlist or --kind placement")
    if selected_kind == "netlist" and selected_family not in {"lpddr", "pcie", "serdes", "pll_clk"}:
        raise AdapterError(
            "Netlist normalization requires the workflow family; pass an explicit --family"
        )
    if selected_topology not in {"pcb", "pkg"}:
        raise AdapterError("Topology is ambiguous; pass --topology pcb or --topology pkg")
    if analysis.table.ambiguous_columns:
        raise AdapterError("Column mapping is ambiguous; refine the selected profile")
    if kind == "auto" and analysis.kind.name in {"unknown", "ambiguous"}:
        raise AdapterError("Input kind is not safe for automatic normalization")
    destination = Path(output).expanduser().resolve()
    if destination.exists() and not force:
        raise AdapterError(f"Output already exists; use --force to replace it: {destination}")
    if destination.suffix.lower() != ".xlsx":
        raise AdapterError("Canonical output must use .xlsx")
    if selected_kind == "netlist":
        rows, ir_rows, statistics = _normalize_netlist(
            analysis, profile, selected_family, selected_topology, groups
        )
        headers = NETLIST_HEADERS
    else:
        rows, statistics = _normalize_placement(analysis, profile)
        ir_rows = []
        headers = PLACEMENT_HEADERS
    _write_workbook(destination, headers, rows, ir_rows)
    report = {
        "schema_version": 1,
        "status": "NORMALIZED",
        "source": {"path": str(analysis.source), "sha256": _sha256(analysis.source)},
        "output": {"path": str(destination), "sha256": _sha256(destination)},
        "profile": {"id": profile["id"], "version": profile["version"], "sha256": profile_sha256},
        "selection": {
            "kind": selected_kind,
            "family": selected_family,
            "topology": selected_topology,
            "groups": groups,
        },
        "source_table": {"sheet": analysis.table.sheet, "header_row": analysis.table.header_row},
        "column_mapping": {
            field: analysis.table.headers[index]
            for field, index in sorted(analysis.table.mapping.items())
        },
        "statistics": statistics,
        "analysis": analysis.as_dict(),
    }
    sidecar = destination.with_suffix(destination.suffix + ".adapter.json")
    sidecar.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return report
