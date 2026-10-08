#!/usr/bin/env python3
"""Generate CHIP/CONNECTOR Placement Excel for SerDes netlists/layouts."""

from __future__ import annotations

import argparse
import csv
import hashlib
import os
import re
import subprocess
import sys
import tempfile
from contextlib import contextmanager
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from artifact_paths import artifact_path
from cadence_runtime import resolve_cadence_export_command as _resolve_cadence_export_command

csv.field_size_limit(64 * 1024 * 1024)

FORMAT_VERSION = "ketupa-serdes-pcb-placement-v3"
PROJECT_ROOT = Path(os.environ["KETUPA_WORKFLOW_INPUT"]).resolve()
DESIGN_SUFFIXES = {".brd", ".mcm", ".sip"}
REPORT_CODES = {"netlist": "net", "component": "cmp"}
REPORT_PREFIXES = {
    "netlist": "SerDesNetListReport",
    "component": "ComponentReport",
}
VARIANT_SUFFIX = re.compile(r"_v(?P<version>[12])$", re.I)
HEADERS = (
    "REFDES",
    "ROLE",
    "SIDE",
    "SYM_MIRROR",
    "SYM_X",
    "SYM_Y",
    "SYM_ROTATE",
    "COMP_PACKAGE",
)

HEADER_COLOR = "17365D"
HEADER_ACCENT = "4472C4"
TEXT_COLOR = "243447"
GRID_COLOR = "D6DEE8"
ROLE_COLORS = {
    "CHIP": "D9EAF7",
    "CONNECTOR": "E4DFEC",
}
ROLE_RANK = {"CHIP": 0, "CONNECTOR": 1}
CHIP_REFDES = re.compile(r"^U[A-Z0-9_.-]*$", re.I)
CONNECTOR_REFDES = re.compile(r"^J[A-Z0-9_.-]*$", re.I)
SIDE_COLORS = {"TOP": "E2F0D9", "BOTTOM": "FFF2CC"}


class TableParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.row = None
        self.cell = None
        self.rows = []

    def handle_starttag(self, tag, attrs):
        del attrs
        tag = tag.casefold()
        if tag == "tr":
            self.row = []
        elif tag in {"td", "th"} and self.row is not None:
            self.cell = []

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)

    def handle_endtag(self, tag):
        tag = tag.casefold()
        if tag in {"td", "th"} and self.cell is not None:
            self.row.append(" ".join("".join(self.cell).replace("\xa0", " ").split()))
            self.cell = None
        elif tag == "tr" and self.row is not None:
            if self.row:
                self.rows.append(self.row)
            self.row = None


def _clean_cell(value):
    return " ".join(str(value).replace("\xa0", " ").split())


def _decode_report(path):
    raw = Path(path).read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("Cannot decode Cadence report: {0}".format(path))


def report_variant(path, text=None):
    path = Path(path)
    text = _decode_report(path) if text is None else text
    content_variant = "v1" if "<table" in text.casefold() else "v2"
    suffix_match = VARIANT_SUFFIX.search(path.stem)
    if suffix_match is not None:
        named_variant = "v" + suffix_match.group("version")
        if named_variant != content_variant:
            raise ValueError(
                "Report suffix/content mismatch for {0}: name={1}, content={2}".format(
                    path, named_variant, content_variant
                )
            )
    return content_variant


def canonical_report_stem(path):
    return VARIANT_SUFFIX.sub("", Path(path).stem)


def canonical_xlsx_path(path, output_path=None):
    path = Path(path).resolve()
    selected = (
        Path(output_path).resolve()
        if output_path is not None
        else path.with_name(canonical_report_stem(path) + ".xlsx")
    )
    if selected.suffix.casefold() != ".xlsx":
        raise ValueError("Output must use .xlsx: {0}".format(selected))
    return selected


def _records_from_rows(rows, required_headers):
    requested = tuple(str(header).strip() for header in required_headers)
    required = {header.casefold() for header in requested}
    for index, row in enumerate(rows):
        normalized = [_clean_cell(value).casefold() for value in row]
        if not required.issubset(set(normalized)):
            continue
        positions = {
            header: normalized.index(header.casefold()) for header in requested
        }
        last_column = max(positions.values())
        records = []
        for values in rows[index + 1 :]:
            cleaned = [_clean_cell(value) for value in values]
            if len(cleaned) <= last_column or not any(cleaned):
                continue
            records.append(
                {header: cleaned[column] for header, column in positions.items()}
            )
        if records:
            return records
    raise ValueError(
        "Cadence report does not contain the required columns: {0}".format(
            ", ".join(requested)
        )
    )


def read_report_table(path, required_headers):
    path = Path(path).resolve()
    if path.suffix.casefold() not in {".htm", ".html"} or not path.is_file():
        raise ValueError("Report must be an existing .htm/.html file: {0}".format(path))
    text = _decode_report(path)
    if report_variant(path, text) == "v1":
        parser = TableParser()
        parser.feed(text)
        rows = parser.rows
    else:
        rows = list(csv.reader(text.splitlines()))
    return _records_from_rows(rows, required_headers)


def resolve_cadence_export_command():
    return _resolve_cadence_export_command()


def default_design_output(project_root, design_path, report_kind):
    return artifact_path(project_root, design_path, report_kind, "serdes")


@contextmanager
def temporary_v2_report(design_path, report_kind, project_root=None):
    del project_root
    design_path = Path(design_path).resolve()
    if design_path.suffix.casefold() not in DESIGN_SUFFIXES or not design_path.is_file():
        raise ValueError(
            "Design must be an existing .brd/.mcm/.sip file: {0}".format(
                design_path
            )
        )
    executable = resolve_cadence_export_command()
    with tempfile.TemporaryDirectory(prefix="ketupa_cadence_report_") as directory:
        report_path = Path(directory) / "{0}_{1}_v2.htm".format(
            REPORT_PREFIXES[report_kind], design_path.stem
        )
        completed = subprocess.run(
            [
                str(executable),
                "-v",
                REPORT_CODES[report_kind],
                str(design_path),
                str(report_path),
            ],
            cwd=directory,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            errors="replace",
            check=False,
        )
        if completed.returncode != 0 or not report_path.is_file():
            raise RuntimeError(
                "Cadence export failed ({0}): {1}".format(
                    completed.returncode,
                    completed.stdout.strip() or "no diagnostic output",
                )
            )
        report_variant(report_path)
        try:
            yield report_path
        finally:
            pass


def read_table(path, required_headers):
    return read_report_table(path, required_headers)


def _darken(color, amount=0.28):
    value = color.lstrip("#")
    rgb = (int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))
    return "".join("{0:02X}".format(round(item * (1 - amount))) for item in rgb)


def style_placement_sheet(sheet, records):
    for cell in sheet[1][:8]:
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_COLOR)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=Side(style="medium", color=HEADER_ACCENT))
    sheet.row_dimensions[1].height = 28
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:H{0}".format(sheet.max_row)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90
    for column, width in zip("ABCDEFGH", (18, 12, 14, 16, 14, 14, 16, 42)):
        sheet.column_dimensions[column].width = width

    previous_role = None
    for row_index, record in enumerate(records, start=2):
        role_color = ROLE_COLORS[record.role]
        role_changed = previous_role != record.role
        for column_index in range(1, 9):
            cell = sheet.cell(row_index, column_index)
            fill_color = (
                SIDE_COLORS[record.side] if column_index in {3, 4} else role_color
            )
            cell.fill = PatternFill("solid", fgColor=fill_color)
            cell.font = Font(
                name="Microsoft YaHei",
                size=9,
                bold=column_index in {1, 2, 3},
                color=TEXT_COLOR,
            )
            cell.alignment = Alignment(
                horizontal="left" if column_index in {1, 8} else "center",
                vertical="center",
            )
            cell.border = Border(
                top=(
                    Side(style="medium", color=_darken(role_color))
                    if role_changed
                    else Side()
                ),
                bottom=Side(style="hair", color=GRID_COLOR),
            )
        sheet.row_dimensions[row_index].height = 20
        previous_role = record.role
    sheet.sheet_properties.tabColor = HEADER_COLOR


@dataclass(frozen=True)
class PlacementRecord:
    refdes: str
    role: str
    side: str
    mirror: str
    x: str
    y: str
    rotate: str
    package: str

    def values(self):
        return (
            self.refdes,
            self.role,
            self.side,
            self.mirror,
            self.x,
            self.y,
            self.rotate,
            self.package,
        )


def natural_key(value):
    return tuple(
        (0, int(part)) if part.isdigit() else (1, part.casefold())
        for part in re.split(r"(\d+)", str(value))
        if part
    )


def refdes_family_key(value):
    """Keep a base RefDes immediately before all of its underscore variants."""
    refdes = str(value).strip()
    base, separator, suffix = refdes.partition("_")
    return (
        natural_key(base),
        1 if separator else 0,
        natural_key(suffix),
        natural_key(refdes),
    )


def split_pin_cell(value):
    return [token for token in re.split(r"[\s,;|]+", str(value or "").strip()) if token]


def refdes_from_pin(pin):
    component, separator, pin_name = str(pin).partition(".")
    if not separator or not component.strip() or not pin_name.strip():
        raise ValueError(
            "Invalid RefDes.Pin token in netlist workbook: {0}".format(pin)
        )
    return component.strip().upper()


def component_role(refdes, pin_column):
    """Require Netlist PIN 1=U* CHIP and PIN 2=J* CONNECTOR."""
    normalized = str(refdes).strip().upper()
    if CHIP_REFDES.fullmatch(normalized):
        if pin_column != 1:
            raise ValueError(
                "U* CHIP must be in Netlist PIN 1, not PIN {0}: {1}".format(
                    pin_column, refdes
                )
            )
        return "CHIP"
    if CONNECTOR_REFDES.fullmatch(normalized):
        if pin_column != 2:
            raise ValueError(
                "J* CONNECTOR must be in Netlist PIN 2, not PIN {0}: {1}".format(
                    pin_column, refdes
                )
            )
        return "CONNECTOR"
    raise ValueError(
        "Unsupported SerDes PCB RefDes role: {0}; expected U* CHIP or J* "
        "CONNECTOR".format(refdes)
    )


def load_netlist_roles(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = next(
        (sheet for sheet in workbook.worksheets if sheet.title.casefold() == "all"),
        workbook.worksheets[0],
    )
    roles = {}
    try:
        rows = worksheet.iter_rows(min_col=1, max_col=6, values_only=True)
        try:
            headers = tuple(str(value or "").strip() for value in next(rows))
        except StopIteration as error:
            raise ValueError("Netlist workbook is empty: {0}".format(path)) from error
        expected = (
            "Net Name",
            "PIN 1",
            "PIN 2",
            "Signal Type",
            "M Group",
            "Chip Partition",
        )
        if headers != expected:
            raise ValueError(
                "Unexpected SerDes Netlist columns: {0}; expected {1}".format(
                    headers, expected
                )
            )
        for row_number, row in enumerate(rows, start=2):
            partition = str(row[5] or "").strip().upper()
            if not re.fullmatch(r"CHIP_[1-5]", partition):
                raise ValueError(
                    "Invalid Chip Partition at Netlist row {0}: {1}".format(
                        row_number, partition or "EMPTY"
                    )
                )
            for pin_column, cell in enumerate(row[1:3], start=1):
                for pin in split_pin_cell(cell):
                    refdes = refdes_from_pin(pin)
                    role = component_role(refdes, pin_column)
                    previous = roles.get(refdes)
                    if previous is not None and previous != role:
                        raise ValueError(
                            "SerDes component appears on both endpoint columns: "
                            "{0}".format(refdes)
                        )
                    roles[refdes] = role
    finally:
        workbook.close()
    if not roles:
        raise ValueError(
            "Netlist workbook does not contain SerDes endpoint pins in columns B/C: {0}".format(
                path
            )
        )
    return roles


def workbook_is_netlist(path):
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = next(
            (sheet for sheet in workbook.worksheets if sheet.title.casefold() == "all"),
            workbook.worksheets[0],
        )
        headers = tuple(
            str(worksheet.cell(1, column).value or "").strip() for column in range(1, 4)
        )
        workbook.close()
        return headers == ("Net Name", "PIN 1", "PIN 2")
    except Exception:
        return False


def name_tokens(path):
    ignored = {"component", "components", "report", "net", "netlist", "placement"}
    return {
        token.casefold()
        for token in re.findall(r"[A-Za-z0-9]+", Path(path).stem)
        if token.casefold() not in ignored
    }


def discover_netlist_workbook(component_path, requested=None):
    if requested is not None:
        selected = Path(requested).resolve()
        if not selected.is_file() or not workbook_is_netlist(selected):
            raise ValueError(
                "Selected file is not a valid netlist workbook: {0}".format(selected)
            )
        return selected
    project_root = Path(os.environ["KETUPA_WORKFLOW_INPUT"]).resolve()
    netlist_directory = project_root / "pcb"
    candidates = [
        path.resolve()
        for path in netlist_directory.glob("*.xls*")
        if not path.name.startswith("~$") and workbook_is_netlist(path)
    ]
    if not candidates:
        raise FileNotFoundError(
            "No netlist workbook was found in {0}".format(netlist_directory)
        )
    component_tokens = name_tokens(component_path)
    scored = sorted(
        candidates,
        key=lambda path: (
            len(component_tokens.intersection(name_tokens(path))),
            path.stat().st_mtime_ns,
            path.name.casefold(),
        ),
        reverse=True,
    )
    if len(scored) > 1:
        first_score = len(component_tokens.intersection(name_tokens(scored[0])))
        second_score = len(component_tokens.intersection(name_tokens(scored[1])))
        if (
            first_score == second_score
            and scored[0].stat().st_mtime_ns == scored[1].stat().st_mtime_ns
        ):
            raise RuntimeError(
                "Cannot select one netlist workbook for {0}: {1}".format(
                    component_path, ", ".join(path.name for path in scored)
                )
            )
    return scored[0]


def load_placements(component_path, netlist_path):
    roles = load_netlist_roles(netlist_path)
    wanted = set(roles)
    records = []
    seen = set()
    unplaced = []
    required = (
        "REFDES",
        "SYM_MIRROR",
        "SYM_X",
        "SYM_Y",
        "SYM_ROTATE",
        "COMP_PACKAGE",
    )
    for row in read_table(component_path, required):
        refdes = row["REFDES"].strip().upper()
        if refdes not in wanted:
            continue
        seen.add(refdes)
        mirror = row["SYM_MIRROR"].strip().upper()
        placement_values = tuple(
            row.get(field, "").strip()
            for field in ("SYM_X", "SYM_Y", "SYM_ROTATE", "SYM_MIRROR")
        )
        if not any(placement_values):
            unplaced.append(refdes)
            continue
        if mirror == "NO":
            side = "TOP"
        elif mirror == "YES":
            side = "BOTTOM"
        else:
            raise ValueError(
                "Unsupported SYM_MIRROR for {0}: {1}".format(refdes, mirror)
            )
        role = roles[refdes]
        records.append(
            PlacementRecord(
                refdes,
                role,
                side,
                mirror,
                row.get("SYM_X", ""),
                row.get("SYM_Y", ""),
                row.get("SYM_ROTATE", ""),
                row.get("COMP_PACKAGE", ""),
            )
        )
    missing = sorted(wanted - seen, key=natural_key)
    if missing:
        raise ValueError(
            "SerDes RefDes values from the netlist are missing from Component report: {0}".format(
                ", ".join(missing)
            )
        )
    if unplaced:
        print(
            "Skipped unplaced SerDes rows (no X/Y/rotation/mirror): {0}".format(
                ", ".join(sorted(unplaced, key=natural_key))
            )
        )
    return sorted(
        records,
        key=lambda row: (ROLE_RANK[row.role], refdes_family_key(row.refdes)),
    )


def records_fingerprint(records):
    digest = hashlib.sha256(FORMAT_VERSION.encode("ascii"))
    for record in records:
        for value in record.values():
            digest.update(str(value).encode("utf-8"))
            digest.update(b"\0")
    return digest.hexdigest().upper()


def output_is_current(path, digest):
    if not path.is_file():
        return False
    try:
        workbook = load_workbook(path, read_only=True)
        current = workbook.properties.keywords == "source-sha256=" + digest
        workbook.close()
        return current
    except Exception:
        return False


def create_workbook(records, digest):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Placement"
    sheet.append(HEADERS)
    for record in records:
        sheet.append(record.values())
    style_placement_sheet(sheet, records)
    workbook.properties.title = "Ketupa SerDes PCB Placement"
    workbook.properties.keywords = "source-sha256=" + digest
    return workbook


def verify_workbook(path, records):
    workbook = load_workbook(path, read_only=True, data_only=True)
    if workbook.sheetnames != ["Placement"]:
        raise AssertionError("Placement workbook must contain one Placement sheet")
    values = list(workbook["Placement"].iter_rows(values_only=True))
    workbook.close()
    normalized = [
        tuple("" if value is None else value for value in row) for row in values[1:]
    ]
    if values[0] != HEADERS or normalized != [record.values() for record in records]:
        raise AssertionError(
            "Placement workbook does not match the source intersection"
        )


def convert(component_path, netlist_path=None, output_path=None, force=False):
    component_path = Path(component_path).resolve()
    if (
        component_path.suffix.casefold() not in {".htm", ".html"}
        or not component_path.is_file()
    ):
        raise ValueError(
            "Component input must be an existing .htm/.html: {0}".format(component_path)
    )
    netlist_path = discover_netlist_workbook(component_path, netlist_path)
    output_path = canonical_xlsx_path(component_path, output_path)
    records = load_placements(component_path, netlist_path)
    digest = records_fingerprint(records)
    if not force and output_is_current(output_path, digest):
        print("Unchanged: {0}".format(output_path))
        return output_path
    temporary = output_path.with_name(".{0}.tmp.xlsx".format(output_path.stem))
    workbook = create_workbook(records, digest)
    try:
        workbook.save(temporary)
        verify_workbook(temporary, records)
        temporary.replace(output_path)
    finally:
        workbook.close()
        if temporary.exists():
            temporary.unlink()
    print("Netlist: {0}".format(netlist_path))
    print("Created: {0}".format(output_path))
    print(
        "Placement rows: {0}; {1}".format(
            len(records),
            ", ".join(
                "{0}={1}".format(role, sum(record.role == role for record in records))
                for role in ROLE_COLORS
            ),
        )
    )
    return output_path


def convert_source(
    component_path,
    netlist_path=None,
    output_path=None,
    force=False,
):
    """Convert v1/v2 component HTML or create both workbooks from BRD/MCM/SIP."""
    component_path = Path(component_path).resolve()
    if component_path.suffix.casefold() in DESIGN_SUFFIXES:
        if netlist_path is None:
            netlist_path = default_design_output(
                PROJECT_ROOT, component_path, "netlist"
            )
            if force or not netlist_path.is_file():
                command = [
                    sys.executable,
                    str(Path(__file__).with_name("serdes_pcb_netlist.py")),
                    str(component_path),
                ]
                if force:
                    command.append("--force")
                completed = subprocess.run(command, check=False)
                if completed.returncode != 0 or not netlist_path.is_file():
                    raise RuntimeError(
                        "Netlist conversion failed for: {0}".format(component_path)
                    )
        output_path = Path(
            output_path
            or default_design_output(PROJECT_ROOT, component_path, "component")
        ).resolve()
        with temporary_v2_report(
            component_path,
            "component",
            PROJECT_ROOT,
        ) as report_path:
            result = convert(report_path, netlist_path, output_path, force)
        print("Removed temporary v2 report.")
        return result
    return convert(component_path, netlist_path, output_path, force)


def discover_component_reports(directory):
    """Return all HTM/HTML reports directly inside *directory*."""
    directory = Path(directory).resolve()
    if not directory.is_dir():
        raise ValueError(
            "Component input directory does not exist: {0}".format(directory)
        )
    return sorted(
        (
            path.resolve()
            for path in directory.iterdir()
            if path.is_file() and path.suffix.casefold() in {".htm", ".html"}
        ),
        key=lambda path: path.name.casefold(),
    )


def convert_directory(
    component_directory, netlist_path=None, output_directory=None, force=False
):
    """Convert every HTM/HTML report directly inside a directory.

    Processing continues after an individual report fails.  The returned
    ``failures`` list contains ``(input_path, exception)`` tuples so callers
    can report a useful batch summary and return a non-zero status.
    """
    component_directory = Path(component_directory).resolve()
    component_paths = discover_component_reports(component_directory)
    if not component_paths:
        raise FileNotFoundError(
            "No .htm or .html component report was found in {0}".format(
                component_directory
            )
        )

    if output_directory is not None:
        output_directory = Path(output_directory).resolve()
        if output_directory.exists() and not output_directory.is_dir():
            raise ValueError(
                "Batch output must be a directory: {0}".format(output_directory)
            )
        output_directory.mkdir(parents=True, exist_ok=True)

    outputs = []
    failures = []
    for component_path in component_paths:
        print("[COMPONENT] {0}".format(component_path.name))
        output_path = None
        if output_directory is not None:
            output_path = output_directory / (
                canonical_report_stem(component_path) + ".xlsx"
            )
        try:
            outputs.append(convert(component_path, netlist_path, output_path, force))
        except Exception as error:  # Continue with the remaining reports.
            failures.append((component_path, error))
            print("[FAILED] {0}: {1}".format(component_path.name, error))

    print(
        "Component reports: {0}  Succeeded: {1}  Failed: {2}".format(
            len(component_paths), len(outputs), len(failures)
        )
    )
    return outputs, failures


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "component",
        nargs="?",
        type=Path,
        default=PROJECT_ROOT,
        help=(
            "one component _v1/_v2 .htm/.html file, one .brd/.mcm/.sip design, "
            "or a directory whose directly contained reports will all be "
            "processed (default: active PCB folder)"
        ),
    )
    parser.add_argument(
        "--netlist",
        type=Path,
        help="one netlist workbook to use for every input; default: auto-match",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        help="output .xlsx for one input, or output directory for batch input",
    )
    parser.add_argument("--force", action="store_true")
    return parser.parse_args()


def main():
    args = parse_args()
    component = args.component.resolve()
    if component.is_dir() or not component.suffix:
        component.mkdir(parents=True, exist_ok=True)
        if not discover_component_reports(component):
            print("No component .htm/.html files found in: {0}".format(component))
            print("Put _v1.htm or _v2.htm files there, then double-click 02 again.")
            return 0
        _, failures = convert_directory(
            component, args.netlist, args.output, args.force
        )
        return 1 if failures else 0
    try:
        convert_source(
            component,
            args.netlist,
            args.output,
            args.force,
        )
    except Exception as error:
        print("[FAILED] {0}: {1}".format(component.name, error))
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
