"""Read Excel/CSV sources and locate tables by semantic column aliases."""

from __future__ import annotations

import csv
from html.parser import HTMLParser
import re
import unicodedata
from pathlib import Path
from typing import Any, Iterable

from .models import AdapterError, TableCandidate


SUPPORTED_SUFFIXES = {".xlsx", ".xlsm", ".csv", ".tsv", ".htm", ".html"}


class _HTMLTableParser(HTMLParser):
    """Read legacy Cadence HTML reports without browser dependencies."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.row: list[str] | None = None
        self.cell: list[str] | None = None
        self.rows: list[tuple[str, ...]] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        if tag.casefold() == "tr":
            self.row = []
        elif tag.casefold() in {"td", "th"} and self.row is not None:
            self.cell = []

    def handle_data(self, data: str) -> None:
        if self.cell is not None:
            self.cell.append(data)

    def handle_endtag(self, tag: str) -> None:
        normalized = tag.casefold()
        if normalized in {"td", "th"} and self.cell is not None:
            assert self.row is not None
            self.row.append(" ".join("".join(self.cell).replace("\xa0", " ").split()))
            self.cell = None
        elif normalized == "tr" and self.row is not None:
            if self.row:
                self.rows.append(tuple(self.row))
            self.row = None


def normalized_label(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).strip().casefold()
    return re.sub(r"[^\w\u4e00-\u9fff]+", "", text)


def _alias_index(profile: dict) -> dict[str, set[str]]:
    return {
        field: {normalized_label(alias) for alias in aliases}
        for field, aliases in profile["columns"].items()
    }


def map_headers(headers: Iterable[Any], profile: dict) -> tuple[dict[str, int], dict[str, tuple[int, ...]]]:
    values = tuple(normalized_label(value) for value in headers)
    aliases = _alias_index(profile)
    mapping: dict[str, int] = {}
    ambiguous: dict[str, tuple[int, ...]] = {}
    for field, accepted in aliases.items():
        matches = tuple(index for index, value in enumerate(values) if value and value in accepted)
        if len(matches) == 1:
            mapping[field] = matches[0]
        elif len(matches) > 1:
            ambiguous[field] = matches
    return mapping, ambiguous


def _trim(row: Iterable[Any]) -> tuple[Any, ...]:
    values = list(row)
    while values and values[-1] is None:
        values.pop()
    return tuple(values)


def _excel_tables(path: Path, profile: dict) -> list[TableCandidate]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise AdapterError("openpyxl is required to read Excel workbooks") from exc
    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
        keep_vba=path.suffix.lower() == ".xlsm",
    )
    candidates: list[TableCandidate] = []
    scan_rows = int(profile.get("table", {}).get("header_scan_rows", 20))
    try:
        for sheet in workbook.worksheets:
            raw = [_trim(row) for row in sheet.iter_rows(values_only=True)]
            best: tuple[int, int, dict[str, int], dict[str, tuple[int, ...]]] | None = None
            for index, row in enumerate(raw[:scan_rows]):
                mapping, ambiguous = map_headers(row, profile)
                rank = len(mapping) * 10 - len(ambiguous)
                item = (rank, -index, mapping, ambiguous)
                if best is None or item[:2] > best[:2]:
                    best = item
            if best is None or best[0] <= 0:
                continue
            header_index = -best[1]
            headers = tuple(str(value or "").strip() for value in raw[header_index])
            rows = tuple(row for row in raw[header_index + 1 :] if any(value not in (None, "") for value in row))
            candidates.append(
                TableCandidate(
                    source=path,
                    sheet=sheet.title,
                    header_row=header_index + 1,
                    headers=headers,
                    rows=rows,
                    mapping=best[2],
                    ambiguous_columns=best[3],
                )
            )
    finally:
        workbook.close()
    return candidates


def _delimited_table(path: Path, profile: dict) -> list[TableCandidate]:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        raw = [tuple(row) for row in csv.reader(stream, delimiter=delimiter)]
    scan_rows = int(profile.get("table", {}).get("header_scan_rows", 20))
    ranked = []
    for index, row in enumerate(raw[:scan_rows]):
        mapping, ambiguous = map_headers(row, profile)
        ranked.append((len(mapping) * 10 - len(ambiguous), -index, mapping, ambiguous))
    if not ranked:
        return []
    best = max(ranked, key=lambda item: item[:2])
    if best[0] <= 0:
        return []
    header_index = -best[1]
    return [
        TableCandidate(
            source=path,
            sheet=path.stem,
            header_row=header_index + 1,
            headers=tuple(raw[header_index]),
            rows=tuple(row for row in raw[header_index + 1 :] if any(str(value).strip() for value in row)),
            mapping=best[2],
            ambiguous_columns=best[3],
        )
    ]


def _decode_report(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise AdapterError(f"Cannot decode Cadence report: {path}")


def _report_table(path: Path, profile: dict) -> list[TableCandidate]:
    """Treat Cadence v1 HTML and v2 CSV-with-.htm-suffix as regular tables."""
    text = _decode_report(path)
    if "<table" in text.casefold():
        parser = _HTMLTableParser()
        parser.feed(text)
        raw = parser.rows
    else:
        raw = [tuple(row) for row in csv.reader(text.splitlines())]
    scan_rows = int(profile.get("table", {}).get("header_scan_rows", 20))
    ranked = []
    for index, row in enumerate(raw[:scan_rows]):
        mapping, ambiguous = map_headers(row, profile)
        ranked.append((len(mapping) * 10 - len(ambiguous), -index, mapping, ambiguous))
    if not ranked:
        return []
    best = max(ranked, key=lambda item: item[:2])
    if best[0] <= 0:
        return []
    header_index = -best[1]
    headers = tuple(str(value or "").strip() for value in raw[header_index])
    return [
        TableCandidate(
            source=path,
            sheet=path.stem,
            header_row=header_index + 1,
            headers=headers,
            rows=tuple(
                row
                for row in raw[header_index + 1 :]
                if any(str(value or "").strip() for value in row)
            ),
            mapping=best[2],
            ambiguous_columns=best[3],
        )
    ]


def read_tables(source: str | Path, profile: dict) -> tuple[TableCandidate, ...]:
    path = Path(source).expanduser().resolve()
    if not path.is_file():
        raise AdapterError(f"Input file not found: {path}")
    if path.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise AdapterError(f"Unsupported table format: {path.suffix}")
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        tables = _excel_tables(path, profile)
    elif path.suffix.lower() in {".htm", ".html"}:
        tables = _report_table(path, profile)
    else:
        tables = _delimited_table(path, profile)
    if not tables:
        raise AdapterError("No recognizable table header was found")
    return tuple(tables)


def cell(row: tuple[Any, ...], index: int | None) -> str:
    if index is None or index >= len(row) or row[index] is None:
        return ""
    return str(row[index]).strip()
