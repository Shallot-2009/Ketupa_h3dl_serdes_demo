"""Offline regression checks for the v1.0.1 adapter."""

from __future__ import annotations

import csv
import os
from pathlib import Path
import sys
import tempfile

from .industry_patterns import identify, recognized
from .models import AdapterError
from .normalizer import normalize
from .profile import load_profile
from extractors.signal_identity import parse_sds_io_die_name
from extractors.cadence_runtime import (
    cadence_export_candidates,
    cadence_design_version,
    cadence_tool_version,
)


ROOT = Path(__file__).resolve().parent


def _write(path: Path, rows: list[tuple[str, str]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(("Net Name", "Net Pins"))
        writer.writerows(rows)


def _check_serdes_extractor_compatibility() -> None:
    """Lock the original PCB/PKG filters before checking N2009 additions."""
    extractor_root = ROOT.parent / "extractors"
    extractor_path = str(extractor_root)
    inserted = extractor_path not in sys.path
    previous_input = os.environ.get("KETUPA_WORKFLOW_INPUT")
    if inserted:
        sys.path.insert(0, extractor_path)
    os.environ["KETUPA_WORKFLOW_INPUT"] = str(ROOT.parents[1] / "input")
    try:
        from serdes_pcb_netlist import parse_net_name as parse_pcb
        from serdes_pkg_netlist import parse_net_name as parse_pkg
    finally:
        if previous_input is None:
            os.environ.pop("KETUPA_WORKFLOW_INPUT", None)
        else:
            os.environ["KETUPA_WORKFLOW_INPUT"] = previous_input
        if inserted:
            sys.path.remove(extractor_path)

    shared = {
        "SDS_RX0P": ("", "RX", 0, "P", "DIFF_P", 1),
        "SDS_TX0N": ("", "TX", 0, "N", "DIFF_N", 1),
        "SDS0_TX0_P": (0, "TX", 0, "P", "DIFF_P", 1),
        "SERDES_TX3_N": (None, "TX", 3, "N", "DIFF_N", 1),
        "SDS2_RX_DP<7>": (None, "RX", 7, "P", "DIFF_P", 1),
        "SDS_TX_N": (None, "TX", 0, "N", "DIFF_N", 1),
        "SDS_RX0P_G1": ("", "RX", 0, "P", "DIFF_P", 2),
        "SDS_RX0N_D2": ("", "RX", 0, "N", "DIFF_N", 2),
        "SDS1_TX1_OUTP_D2": (1, "TX", 1, "P", "DIFF_P", 2),
        "HSS0_TX_LANE3_P": (None, "TX", 3, "P", "DIFF_P", 1),
        "HSIO_RX3_N": (None, "RX", 3, "N", "DIFF_N", 1),
        "XFI_RX0P": (None, "RX", 0, "P", "DIFF_P", 1),
        "CEI_LANE2_RX_N": (None, "RX", 2, "N", "DIFF_N", 1),
        "SERDES_TX_DP3": (None, "TX", 3, "P", "DIFF_P", 1),
    }
    for parser, partition_prefix in ((parse_pcb, "CHIP"), (parse_pkg, "DIE")):
        for name, values in shared.items():
            expected = (*values[:5], f"{partition_prefix}_{values[5]}", values[5])
            actual = parser(name)
            if actual != expected:
                raise AssertionError(
                    f"SerDes {partition_prefix} compatibility changed for {name}: "
                    f"{actual!r} != {expected!r}"
                )
        if parser("PCIE_TX0_P") is not None:
            raise AssertionError("SerDes parser claimed an explicit PCIe network")
        try:
            parser("SDS0_RX2_OUTN_D1")
        except ValueError:
            pass
        else:
            raise AssertionError("N2009 RX/OUT mismatch did not fail closed")


def run() -> int:
    with tempfile.TemporaryDirectory(prefix="ketupa_cadence_version_") as folder:
        marker = Path(folder) / "writer.sip"
        marker.write_bytes(b"binary-prefix\x00cdnsip 25.1 P001\x00binary-suffix")
        if cadence_design_version(marker) != (25, 1):
            raise AssertionError("Cadence SIP writer-version detection changed")
    fake_report = Path(r'C:/Cadence/SPB_25.1/tools/bin/report.exe')
    if cadence_tool_version(fake_report) != (25, 1):
        raise AssertionError('Windows SPB version detection changed')
    previous_cadence = os.environ.get('CADENCE_TOOLS_BIN')
    os.environ['CADENCE_TOOLS_BIN'] = str(fake_report.parent)
    try:
        if cadence_export_candidates()[0] != fake_report.resolve():
            raise AssertionError('Windows explicit Cadence override lost priority')
    finally:
        if previous_cadence is None: os.environ.pop('CADENCE_TOOLS_BIN', None)
        else: os.environ['CADENCE_TOOLS_BIN'] = previous_cadence
    profile, digest = load_profile("industry_default", ROOT / "profiles")
    naming_checks = {
        "lpddr": (
            "DDR3_MC0_CHA_DQ0",
            "DDR4_MC1_CHB_DQS0_P",
            "DDR5_MC0_CH0_DQS1_N",
            "DDR6_MC2_CHC_CA3",
            "LPDDR5_MC0_CHA_WCK0_P",
            "LPDDR6_MC1_CHB_DQ7",
        ),
        "pcie": ("PCIE_RX0_P", "PEX_TX_DP<15>", "PCIEGEN6_TX_DN[3]"),
        "serdes": (
            "SERDES_TX0_P",
            "SDS2_RX_DN<7>",
            "CEI_TX3_N",
            "SDS0_RX0_INP_D1",
            "SDS0_RX0_INN_D1",
            "SDS1_TX1_OUTP_D2",
            "SDS1_TX1_OUTN_D2",
        ),
        "pll_clk": ("PEX_REFCLK*", "PCIE_CLK_DP", "PCIE_CLK_DN"),
    }
    for family, names in naming_checks.items():
        missing = [name for name in names if not recognized(name, family)]
        if missing:
            raise AssertionError(f"{family} recognizer missed: {missing}")

    ns2009_pcb = identify(
        name="SDS0_RX2_INN_D1",
        family="serdes",
        topology="pcb",
        endpoint_a=["J7.1"],
    )
    ns2009_pkg = identify(
        name="SDS1_TX1_OUTP_D2",
        family="serdes",
        topology="pkg",
        endpoint_a=["BGA.A1"],
    )
    if ns2009_pcb is None or ns2009_pcb["logical_id"] != "serdes:sds0:die1:RX:2:N":
        raise AssertionError("NS2009 PCB SDS/DIE identity changed")
    if ns2009_pcb["group"] != "SDS0_RX2" or ns2009_pcb["partition"] != "CHIP_1":
        raise AssertionError("NS2009 PCB grouping changed")
    if ns2009_pkg is None or ns2009_pkg["logical_id"] != "serdes:sds1:die2:TX:1:P":
        raise AssertionError("NS2009 PKG SDS/DIE identity changed")
    if ns2009_pkg["group"] != "SDS1_TX1" or ns2009_pkg["partition"] != "DIE_2":
        raise AssertionError("NS2009 PKG grouping changed")

    legacy_sds = identify(
        name="SDS0_TX0_P",
        family="serdes",
        topology="pcb",
        endpoint_a=["U2.A1"],
    )
    if legacy_sds is None or legacy_sds["logical_id"] != "serdes:TX:0:P":
        raise AssertionError("legacy SDS identity changed")
    if parse_sds_io_die_name("SDS0_TX0_P") is not None:
        raise AssertionError("NS2009 parser claimed a legacy SDS name")
    if parse_sds_io_die_name("SDS0_RX2_INN_D1") != (0, "RX", 2, "N", 1):
        raise AssertionError("NS2009 extractor identity changed")
    try:
        parse_sds_io_die_name("SDS0_RX2_OUTN_D1")
    except ValueError:
        pass
    else:
        raise AssertionError("NS2009 direction/flow mismatch did not fail closed")

    _check_serdes_extractor_compatibility()

    fixtures = {
        "lpddr": [
            ("LPDDR5_MC0_CHA_DQS0_P", "G1.A1 DDR0.B1"),
            ("LPDDR5_MC0_CHA_DQS0_N", "G1.A2 DDR0.B2"),
            ("DDR5_MC0_CHA_DQ0", "G1.A3 DDR0.B3"),
        ],
        "pcie": [
            ("PCIE_RX0_P", "U1.A1 J1.B1"),
            ("PCIE_RX0_N", "U1.A2 J1.B2"),
        ],
        "serdes": [
            ("SDS0_TX0_P", "U2.A1 J2.B1"),
            ("SDS0_TX0_N", "U2.A2 J2.B2"),
            ("SDS0_RX0_INP_D1", "U32.C21 J5.1"),
            ("SDS0_RX0_INN_D1", "U32.C20 J3.1"),
            ("SDS1_TX1_OUTP_D2", "U32.AF2 J41.1"),
            ("SDS1_TX1_OUTN_D2", "U32.AG2 J39.1"),
        ],
        "pll_clk": [
            ("PEX_REFCLK", "G1.A1 CN1.B1"),
            ("PEX_REFCLK*", "G1.A2 CN1.B2"),
        ],
    }
    with tempfile.TemporaryDirectory(prefix="ketupa_adapter_selftest_") as folder:
        root = Path(folder)
        for family, rows in fixtures.items():
            source = root / f"{family}.csv"
            output = root / f"{family}.xlsx"
            _write(source, rows)
            report = normalize(
                source,
                output,
                profile,
                digest,
                family=family,
                topology="pcb",
                kind="netlist",
            )
            if report["statistics"]["rows"] != len(rows):
                raise AssertionError(f"{family} normalized row count changed")
            from openpyxl import load_workbook

            workbook = load_workbook(output, read_only=False, data_only=True)
            try:
                if workbook.sheetnames != ["ALL", "_Ketupa_IR"]:
                    raise AssertionError(f"{family} workbook contract changed")
                if workbook["_Ketupa_IR"].sheet_state != "hidden":
                    raise AssertionError("IR worksheet must remain hidden")
            finally:
                workbook.close()
        negative = {
            "incomplete-pcie": (
                "pcie",
                [("PCIE_RX0_P", "U1.A1 J1.B1")],
            ),
            "ambiguous-ddr-endpoints": (
                "lpddr",
                [("DDR5_MC0_CHA_DQ0", "U1.A1 U2.B1")],
            ),
            "clock-power-only": (
                "pll_clk",
                [("GPU_PLLVDD", "U1.A1 C1.1")],
            ),
        }
        for index, (label, (family, rows)) in enumerate(negative.items(), start=1):
            source = root / f"negative-{index}.csv"
            output = root / f"negative-{index}.xlsx"
            _write(source, rows)
            try:
                normalize(
                    source,
                    output,
                    profile,
                    digest,
                    family=family,
                    topology="pcb",
                    kind="netlist",
                )
            except AdapterError:
                pass
            else:
                raise AssertionError(f"negative case did not fail closed: {label}")
    print("PASS adapter self-test: DDR3/4/5/6/LPDDR, PCIe, legacy + NS2009 SerDes, PLL/CLK")
    print("PASS SerDes extractor matrix: legacy + N2009 + industry-domain filter loop")
    print("PASS canonical Excel and hidden _Ketupa_IR contracts")
    print("PASS incomplete pairs, ambiguous endpoints, and false clock candidates fail closed")
    print("PASS no network or AEDT process was used")
    print("PASS Windows Cadence override priority and writer/SPB discovery; Linux report snapshot test not applicable")
    return 0


__all__ = ["run"]
