#!/usr/bin/env python3
"""Convert PCIe PCB v1/v2 netlist reports or BRD/MCM/SIP designs into Excel."""

from __future__ import annotations

import argparse
import csv
import hashlib
import os
import re
import subprocess
import tempfile
from collections import Counter
from contextlib import contextmanager
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from artifact_paths import artifact_path, NO_SIGNALS, NoMatchingSignals, extractor_exit
from cadence_runtime import resolve_cadence_export_command as _resolve_cadence_export_command, run_cadence_export
from signal_identity import accepts_signal_family

csv.field_size_limit(64 * 1024 * 1024)

FORMAT_VERSION = "ketupa-pcie-netlist-v2"
PROJECT_ROOT = Path(os.environ["KETUPA_WORKFLOW_INPUT"]).resolve()
DESIGN_SUFFIXES = {".brd", ".mcm", ".sip"}
REPORT_CODES = {"netlist": "net", "component": "cmp"}
REPORT_PREFIXES = {
    "netlist": "PCIeNetListReport",
    "component": "ComponentReport",
}
VARIANT_SUFFIX = re.compile(r"_v(?P<version>[12])$", re.I)
PCIE_NET_PATTERNS = (
    # TX0_P, TX_0_N, PCIE_GPU0_TX7_DP
    re.compile(
        r"(?:^|[_\-.])(?P<direction>TX|RX)[_\-.]?(?P<lane>\d{1,2})"
        r"[_\-.]+D?(?P<polarity>P|N)$",
        re.I,
    ),
    # Cadence bus form: PCIE_GPU0_RX_DP<0>, RX_DN[15]
    re.compile(
        r"(?:^|[_\-.])(?P<direction>TX|RX)[_\-.]+D?(?P<polarity>P|N)"
        r"\s*[<\[(](?P<lane>\d{1,2})[>\])]$",
        re.I,
    ),
    # AC-coupled segment: PCIE_TX_AC_DP<0>, PCIE_TX_CAP_DN[15].
    # Qualifier tokens between TX/RX and DP/DN are intentionally accepted.
    re.compile(
        r"(?:^|[_\-.])(?P<direction>TX|RX)"
        r"(?:[_\-.]+[A-Z0-9]+?)*?[_\-.]+D?(?P<polarity>P|N)"
        r"\s*[<\[(](?P<lane>\d{1,2})[>\])]$",
        re.I,
    ),
)
PCIE_DOMAIN = re.compile(
    r"(?:^|[_\-.])(?P<name>(?:GPU|CPU|ROOT|PORT|PCIE)(?P<index>\d+))"
    r"(?=[_\-.]+(?:TX|RX))",
    re.I,
)
DIRECTION_RANK = {"RX": 0, "TX": 1}
POLARITY_RANK = {"P": 0, "N": 1}
CHIP_SUFFIX = re.compile(r"(?:_|-|\.)G(?P<index>\d*)$", re.I)
PIN = re.compile(r"^(?P<refdes>[^.\s]+)\.(?P<pin>\S+)$")
HEADERS = ("Net Name", "PIN 1", "PIN 2", "Signal Type", "M Group", "Chip Partition")

HEADER_COLOR = "17365D"
HEADER_ACCENT = "4472C4"
TEXT_COLOR = "243447"
GRID_COLOR = "D6DEE8"
BLOCK_COLORS = (
    "D9EAF7",
    "E2F0D9",
    "FCE4D6",
    "E4DFEC",
    "FFF2CC",
    "DDEBF7",
    "EADCF8",
    "F4CCCC",
    "D9EAD3",
    "D0E0E3",
    "FCE5CD",
    "DDEBF1",
)
CHIP_COLORS = (
    "E2F0D9",
    "FCE4D6",
    "DDEBF7",
    "E4DFEC",
    "FFF2CC",
    "D9EAD3",
)
CHANNEL_ACCENTS = ("4472C4", "70AD47", "ED7D31", "A64D79", "5B9BD5")


class TableParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.row = None
        self.cell = None
        self.rows = []

    def handle_starttag(self, tag, attrs):
        del attrs
        tag = tag.casefold()
        if tag == "tr":
            self.row = []
        elif tag in {"td", "th"} and self.row is not None:
            self.cell = []

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)

    def handle_endtag(self, tag):
        tag = tag.casefold()
        if tag in {"td", "th"} and self.cell is not None:
            self.row.append(" ".join("".join(self.cell).replace("\xa0", " ").split()))
            self.cell = None
        elif tag == "tr" and self.row is not None:
            if self.row:
                self.rows.append(self.row)
            self.row = None


def _clean_cell(value):
    return " ".join(str(value).replace("\xa0", " ").split())


def _decode_report(path):
    raw = Path(path).read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("Cannot decode Cadence report: {0}".format(path))


def report_variant(path, text=None):
    path = Path(path)
    text = _decode_report(path) if text is None else text
    content_variant = "v1" if "<table" in text.casefold() else "v2"
    suffix_match = VARIANT_SUFFIX.search(path.stem)
    if suffix_match is not None:
        named_variant = "v" + suffix_match.group("version")
        if named_variant != content_variant:
            raise ValueError(
                "Report suffix/content mismatch for {0}: name={1}, content={2}".format(
                    path, named_variant, content_variant
                )
            )
    return content_variant


def canonical_report_stem(path):
    return VARIANT_SUFFIX.sub("", Path(path).stem)


def canonical_xlsx_path(path, output_path=None):
    path = Path(path).resolve()
    selected = (
        Path(output_path).resolve()
        if output_path is not None
        else path.with_name(canonical_report_stem(path) + ".xlsx")
    )
    if selected.suffix.casefold() != ".xlsx":
        raise ValueError("Output must use .xlsx: {0}".format(selected))
    return selected


def _records_from_rows(rows, required_headers):
    requested = tuple(str(header).strip() for header in required_headers)
    required = {header.casefold() for header in requested}
    for index, row in enumerate(rows):
        normalized = [_clean_cell(value).casefold() for value in row]
        if not required.issubset(set(normalized)):
            continue
        positions = {
            header: normalized.index(header.casefold()) for header in requested
        }
        last_column = max(positions.values())
        records = []
        for values in rows[index + 1 :]:
            cleaned = [_clean_cell(value) for value in values]
            if len(cleaned) <= last_column or not any(cleaned):
                continue
            records.append(
                {header: cleaned[column] for header, column in positions.items()}
            )
        if records:
            return records
    raise ValueError(
        "Cadence report does not contain the required columns: {0}".format(
            ", ".join(requested)
        )
    )


def read_report_table(path, required_headers):
    path = Path(path).resolve()
    if path.suffix.casefold() not in {".htm", ".html"} or not path.is_file():
        raise ValueError("Report must be an existing .htm/.html file: {0}".format(path))
    text = _decode_report(path)
    if report_variant(path, text) == "v1":
        parser = TableParser()
        parser.feed(text)
        rows = parser.rows
    else:
        rows = list(csv.reader(text.splitlines()))
    return _records_from_rows(rows, required_headers)


def resolve_cadence_export_command():
    return _resolve_cadence_export_command()


def default_design_output(project_root, design_path, report_kind):
    return artifact_path(project_root, design_path, report_kind, "pcie")


@contextmanager
def temporary_v2_report(design_path, report_kind, project_root=None):
    del project_root
    design_path = Path(design_path).resolve()
    if design_path.suffix.casefold() not in DESIGN_SUFFIXES or not design_path.is_file():
        raise ValueError(
            "Design must be an existing .brd/.mcm/.sip file: {0}".format(
                design_path
            )
        )
    with tempfile.TemporaryDirectory(prefix="ketupa_cadence_report_") as directory:
        report_path = Path(directory) / "{0}_{1}_v2.htm".format(
            REPORT_PREFIXES[report_kind], design_path.stem
        )
        run_cadence_export(REPORT_CODES[report_kind], design_path, report_path)
        report_variant(report_path)
        try:
            yield report_path
        finally:
            pass


def read_table(path, required_headers):
    return read_report_table(path, required_headers)


def _rgb(value):
    value = value.lstrip("#")
    return tuple(int(value[index : index + 2], 16) for index in (0, 2, 4))


def _hex(values):
    return "".join("{0:02X}".format(max(0, min(255, value))) for value in values)


def _darken(color, amount=0.28):
    return _hex(tuple(round(value * (1 - amount)) for value in _rgb(color)))


def _channel_color(channel):
    number = sum(ord(character) for character in str(channel).upper())
    return CHANNEL_ACCENTS[number % len(CHANNEL_ACCENTS)]


def color_block_key(record):
    net = record.net
    return net.chip, net.direction, net.lane


def style_netlist_sheet(sheet, records):
    for cell in sheet[1][:6]:
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_COLOR)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=Side(style="medium", color=HEADER_ACCENT))
    sheet.row_dimensions[1].height = 28
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:F{0}".format(sheet.max_row)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90
    for column, width in zip("ABCDEF", (34, 30, 44, 14, 12, 18)):
        sheet.column_dimensions[column].width = width

    block_keys = list(dict.fromkeys(color_block_key(record) for record in records))
    block_colors = {
        key: BLOCK_COLORS[index % len(BLOCK_COLORS)]
        for index, key in enumerate(block_keys)
    }
    previous = None
    for row_index, record in enumerate(records, start=2):
        net = record.net
        block_color = block_colors[color_block_key(record)]
        chip_color = CHIP_COLORS[(net.chip - 1) % len(CHIP_COLORS)]
        channel_color = _channel_color(net.direction)
        chip_changed = previous is None or previous.net.chip != net.chip
        group_changed = chip_changed or previous.net.lane != net.lane
        block_changed = group_changed or previous.net.direction != net.direction
        next_record = records[row_index - 1] if row_index - 1 < len(records) else None
        block_ends = next_record is None or color_block_key(
            next_record
        ) != color_block_key(record)
        top_color = _darken(chip_color) if chip_changed else _darken(block_color)
        for column_index in range(1, 7):
            cell = sheet.cell(row_index, column_index)
            cell.font = Font(
                name="Microsoft YaHei",
                size=9,
                color=TEXT_COLOR,
                bold=column_index in {1, 5, 6},
            )
            cell.fill = PatternFill(
                "solid", fgColor=chip_color if column_index == 6 else block_color
            )
            cell.alignment = Alignment(
                horizontal="center" if column_index >= 4 else "left",
                vertical="center",
            )
            cell.border = Border(
                left=(
                    Side(style="medium", color=channel_color)
                    if column_index == 1
                    else Side()
                ),
                top=(
                    Side(
                        style="medium",
                        color=top_color
                        if chip_changed or group_changed
                        else channel_color,
                    )
                    if block_changed
                    else Side()
                ),
                bottom=Side(
                    style="medium" if block_ends else "hair",
                    color=_darken(block_color) if block_ends else GRID_COLOR,
                ),
            )
        sheet.row_dimensions[row_index].height = 19
        previous = record
    sheet.sheet_properties.tabColor = (
        HEADER_COLOR
        if sheet.title == "ALL"
        else _darken(BLOCK_COLORS[records[0].net.lane % len(BLOCK_COLORS)], 0.15)
    )


@dataclass(frozen=True)
class NetRecord:
    name: str
    pins: tuple[str, ...]
    direction: str
    lane: int
    polarity: str
    signal_type: str
    partition: str
    chip: int


@dataclass(frozen=True)
class ExcelRecord:
    net: NetRecord
    pin1: str
    pin2: str

    def values(self):
        return (
            self.net.name,
            self.pin1,
            self.pin2,
            self.net.signal_type,
            "{0}{1}".format(self.net.direction, self.net.lane),
            self.net.partition,
        )


def natural_key(value):
    return tuple(
        (0, int(part)) if part.isdigit() else (1, part.casefold())
        for part in re.split(r"(\d+)", str(value))
        if part
    )


def refdes(pin):
    match = PIN.fullmatch(pin)
    if not match:
        raise ValueError("Invalid RefDes.Pin token: {0}".format(pin))
    return match.group("refdes")


def is_tx_capacitor(component):
    """Return True for a conventional capacitor RefDes such as C51 or C51_1."""
    return re.fullmatch(r"C\d+(?:[_\-.][A-Z0-9]+)*", component, re.I) is not None


def parse_net_name(name):
    clean = name.strip()
    if not accepts_signal_family(clean, "pcie"):
        return None
    suffix = CHIP_SUFFIX.search(clean)
    chip = 1
    classified = clean
    if suffix:
        suffix_index = int(suffix.group("index") or "1")
        chip = suffix_index + 1
        classified = clean[: suffix.start()]
    match = None
    for pattern in PCIE_NET_PATTERNS:
        match = pattern.search(classified)
        if match is not None:
            break
    if match is None:
        return None
    lane = int(match.group("lane"))
    if lane > 15:
        return None
    direction = match.group("direction").upper()
    polarity = match.group("polarity").upper()
    domain_match = PCIE_DOMAIN.search(classified)
    if domain_match is not None:
        domain_chip = int(domain_match.group("index")) + 1
        if suffix and domain_chip != chip:
            raise ValueError(
                "Conflicting PCIe domain and chip suffix in network: {0}".format(
                    name
                )
            )
        chip = domain_chip
        partition = domain_match.group("name").upper()
    else:
        partition = "CHIP_{0}".format(chip)
    return (
        direction,
        lane,
        polarity,
        "DIFF_{0}".format(polarity),
        partition,
        chip,
    )


def explicit_chip_number(value):
    match = CHIP_SUFFIX.search(str(value).strip())
    return None if match is None else int(match.group("index") or "1") + 1


def parse_group_selection(value):
    text = str(value).strip().casefold()
    if text in {"", "auto", "all"}:
        return None
    selected = set()
    for token in text.split(","):
        token = token.strip().removeprefix("lane")
        if "-" in token:
            start, end = token.split("-", 1)
            start_number = int(start.removeprefix("lane"))
            end_number = int(end.removeprefix("lane"))
            if start_number > end_number:
                raise ValueError("Descending PCIe lane range: {0}".format(token))
            selected.update(range(start_number, end_number + 1))
        else:
            selected.add(int(token))
    invalid = sorted(lane for lane in selected if lane < 0 or lane > 15)
    if invalid:
        raise ValueError("PCIe lanes must be in the range 0..15: {0}".format(invalid))
    return selected


def load_net_records(path, selected_groups=None):
    raw_records = []
    for row in read_table(path, ("Net Name", "Net Pins")):
        name = row["Net Name"].strip()
        pins = tuple(
            token for token in re.split(r"[\s,;|]+", row["Net Pins"].strip()) if token
        )
        raw_records.append((name, pins))

    # Classify standard names first, then follow the same C* device to the
    # second TX net. This preserves AC-coupled TX segments even when the
    # connector-side net uses a board-specific name without TX/RX tokens.
    classifications = {}
    capacitor_to_rows = {}
    for row_index, (name, pins) in enumerate(raw_records):
        parsed = parse_net_name(name)
        if parsed is not None:
            classifications[row_index] = parsed
        for pin in pins:
            component = refdes(pin)
            if is_tx_capacitor(component):
                capacitor_to_rows.setdefault(component.casefold(), set()).add(row_index)

    for row_index, parsed in list(classifications.items()):
        direction, lane, polarity, category, partition, chip = parsed
        if direction != "TX":
            continue
        for pin in raw_records[row_index][1]:
            component = refdes(pin)
            if not is_tx_capacitor(component):
                continue
            for linked_index in capacitor_to_rows.get(component.casefold(), ()):
                if linked_index == row_index:
                    continue
                inherited = (direction, lane, polarity, category, partition, chip)
                existing = classifications.get(linked_index)
                if existing is not None and existing != inherited:
                    linked_name = raw_records[linked_index][0]
                    same_lane = existing[:4] == inherited[:4]
                    linked_has_partition = (
                        PCIE_DOMAIN.search(linked_name) is not None
                        or CHIP_SUFFIX.search(linked_name) is not None
                    )
                    if same_lane and not linked_has_partition:
                        # A generic connector-side name inherits GPU/CPU/chip
                        # ownership from the explicitly partitioned host side.
                        existing = inherited
                    else:
                        raise ValueError(
                            "Conflicting PCIe classification across capacitor {0}: {1} / {2}".format(
                                component,
                                raw_records[row_index][0],
                                linked_name,
                            )
                        )
                classifications[linked_index] = inherited

    records = []
    seen = set()
    skipped_single_pin = []
    for row_index, (name, pins) in enumerate(raw_records):
        parsed = classifications.get(row_index)
        if parsed is None:
            continue
        direction, lane, polarity, category, partition, chip = parsed
        if selected_groups is not None and lane not in selected_groups:
            continue
        key = name.casefold()
        if key in seen:
            raise ValueError("Duplicate PCIe network: {0}".format(name))
        seen.add(key)
        if len(pins) < 2:
            skipped_single_pin.append(name)
            continue
        mismatched = []
        for pin in pins:
            component = refdes(pin)
            component_chip = explicit_chip_number(component)
            if component_chip is not None and component_chip != chip:
                mismatched.append(component)
        mismatched = sorted(set(mismatched), key=natural_key)
        if mismatched:
            raise ValueError(
                "Network/RefDes chip partition mismatch for {0}: {1}".format(
                    name, ", ".join(mismatched)
                )
            )
        records.append(
            NetRecord(
                name,
                pins,
                direction,
                lane,
                polarity,
                category,
                partition,
                chip,
            )
        )
    if not records:
        raise NoMatchingSignals("No PCIe TX/RX lane networks were found in {0}".format(path))
    if skipped_single_pin:
        print(
            "Skipped single-pin package nets: {0} ({1})".format(
                len(skipped_single_pin), ", ".join(skipped_single_pin)
            )
        )
    validate_differential_pairs(records)
    return records


def validate_differential_pairs(records):
    """Require balanced P/N segments for every selected PCIe lane."""
    pairs = {}
    for record in records:
        key = (record.partition, record.direction, record.lane)
        pairs.setdefault(key, []).append(record.polarity)
    invalid = [
        "{0}:{1}{2}={3}".format(partition, direction, lane, sorted(polarities))
        for (partition, direction, lane), polarities in sorted(pairs.items())
        if not polarities
        or polarities.count("P") != polarities.count("N")
    ]
    if invalid:
        raise ValueError(
            "Every PCIe lane must contain the same number of P and N segments: {0}".format(
                "; ".join(invalid)
            )
        )


def endpoint_rank(component):
    """Order a PCIe segment as host -> TX capacitor -> connector."""
    if is_tx_capacitor(component):
        return 1
    if re.fullmatch(r"J\d+(?:[_\-.][A-Z0-9]+)*", component, re.I):
        return 2
    return 0


def build_excel_records(records):
    output = []
    for record in records:
        components = sorted(
            dict.fromkeys(refdes(pin) for pin in record.pins),
            key=lambda name: (endpoint_rank(name), natural_key(name)),
        )
        if len(components) < 2:
            raise ValueError(
                "Cannot split the two PCIe segment endpoints for {0}".format(record.name)
            )
        first_component = components[0]
        first_pins = [pin for pin in record.pins if refdes(pin) == first_component]
        target_pins = [pin for pin in record.pins if refdes(pin) != first_component]
        output.append(
            ExcelRecord(record, " ".join(first_pins), " ".join(target_pins))
        )
    return output


def row_sort_key(record):
    net = record.net
    return (
        net.chip,
        DIRECTION_RANK[net.direction],
        net.lane,
        POLARITY_RANK[net.polarity],
        natural_key(net.name),
    )


def records_fingerprint(records):
    digest = hashlib.sha256()
    digest.update(FORMAT_VERSION.encode("ascii"))
    for record in sorted(records, key=row_sort_key):
        for value in record.values():
            digest.update(str(value).encode("utf-8"))
            digest.update(b"\0")
    return digest.hexdigest().upper()


def output_is_current(path, fingerprint):
    if not path.is_file():
        return False
    try:
        workbook = load_workbook(path, read_only=True)
        current = workbook.properties.keywords == "source-sha256=" + fingerprint
        workbook.close()
        return current
    except Exception:
        return False


def write_sheet(workbook, title, records):
    sheet = workbook.create_sheet(title)
    sheet.append(HEADERS)
    for record in records:
        sheet.append(record.values())
    style_netlist_sheet(sheet, records)


def create_workbook(records, fingerprint):
    workbook = Workbook()
    workbook.remove(workbook.active)
    ordered = sorted(records, key=row_sort_key)
    write_sheet(workbook, "ALL", ordered)
    write_sheet(workbook, "RX", [row for row in ordered if row.net.direction == "RX"])
    write_sheet(workbook, "TX", [row for row in ordered if row.net.direction == "TX"])
    workbook.properties.title = "Ketupa PCIe Differential Netlist Classification"
    workbook.properties.keywords = "source-sha256=" + fingerprint
    return workbook


def verify_workbook(path, records):
    workbook = load_workbook(path, read_only=False, data_only=True)
    expected_sheets = ["ALL", "RX", "TX"]
    if workbook.sheetnames != expected_sheets:
        raise AssertionError(
            "Unexpected worksheet order: {0}".format(workbook.sheetnames)
        )
    all_rows = list(workbook["ALL"].iter_rows(values_only=True))
    if all_rows[0] != HEADERS:
        raise AssertionError("The six-column workbook contract changed")
    expected = [row.values() for row in sorted(records, key=row_sort_key)]
    if all_rows[1:] != expected:
        raise AssertionError("ALL worksheet rows do not match the parsed report")
    for direction_sheet in workbook.sheetnames[1:]:
        actual = list(workbook[direction_sheet].iter_rows(min_row=2, values_only=True))
        wanted = [
            row.values()
            for row in sorted(records, key=row_sort_key)
            if row.net.direction == direction_sheet
        ]
        if actual != wanted:
            raise AssertionError(
                "Worksheet {0} is not grouped correctly".format(direction_sheet)
            )
    ordered = sorted(records, key=row_sort_key)
    for sheet_name in workbook.sheetnames:
        sheet_records = (
            ordered
            if sheet_name == "ALL"
            else [row for row in ordered if row.net.direction == sheet_name]
        )
        blocks = []
        for record in sheet_records:
            key = color_block_key(record)
            if not blocks or blocks[-1][0] != key:
                blocks.append((key, []))
            blocks[-1][1].append(record)
        row_number = 2
        previous_color = None
        for key, block_records in blocks:
            start_row = row_number
            end_row = start_row + len(block_records) - 1
            colors = {
                workbook[sheet_name].cell(row, column).fill.fgColor.rgb
                for row in range(start_row, end_row + 1)
                for column in range(1, 6)
            }
            if len(colors) != 1:
                raise AssertionError(
                    "Worksheet {0} block {1} does not use one background color".format(
                        sheet_name, key
                    )
                )
            color = next(iter(colors))
            if color == previous_color:
                raise AssertionError(
                    "Worksheet {0} adjacent blocks share one color: {1}".format(
                        sheet_name, key
                    )
                )
            if workbook[sheet_name].cell(start_row, 1).border.top.style != "medium":
                raise AssertionError(
                    "Worksheet {0} block {1} has no top separator".format(
                        sheet_name, key
                    )
                )
            if workbook[sheet_name].cell(end_row, 1).border.bottom.style != "medium":
                raise AssertionError(
                    "Worksheet {0} block {1} has no bottom separator".format(
                        sheet_name, key
                    )
                )
            previous_color = color
            row_number = end_row + 1
    workbook.close()


def _legacy_convert(input_path, output_path=None, groups="auto", force=False):
    input_path = Path(input_path).resolve()
    if (
        input_path.suffix.casefold() not in {".htm", ".html"}
        or not input_path.is_file()
    ):
        raise ValueError(
            "Input must be an existing .htm or .html file: {0}".format(input_path)
        )
    output_path = canonical_xlsx_path(input_path, output_path)
    records = build_excel_records(
        load_net_records(input_path, parse_group_selection(groups))
    )
    fingerprint = records_fingerprint(records)
    if not force and output_is_current(output_path, fingerprint):
        print("Unchanged: {0}".format(output_path))
        return output_path
    temporary = output_path.with_name(".{0}.tmp.xlsx".format(output_path.stem))
    workbook = create_workbook(records, fingerprint)
    try:
        workbook.save(temporary)
        verify_workbook(temporary, records)
        temporary.replace(output_path)
    finally:
        workbook.close()
        if temporary.exists():
            temporary.unlink()
    counts = Counter(record.net.direction for record in records)
    print("Created: {0}".format(output_path))
    print(
        "Networks: {0}; sheets: {1}".format(
            len(records),
            "ALL, RX, TX",
        )
    )
    print(
        "Directions: "
        + ", ".join(
            "{0}={1}".format(direction, counts[direction])
            for direction in ("RX", "TX")
        )
    )
    return output_path


def convert(input_path, output_path=None, groups="auto", force=False):
    """Keep the qualified legacy path and use the approved adapter on rejection."""
    input_path = Path(input_path).resolve()
    selected_output = canonical_xlsx_path(input_path, output_path)
    from script.data_adapter.runtime import convert_with_fallback

    return convert_with_fallback(
        lambda: _legacy_convert(input_path, selected_output, groups, force),
        input_path,
        selected_output,
        family="pcie",
        topology="pcb",
        groups=groups,
        force=force,
    )


def convert_source(
    input_path,
    output_path=None,
    groups="auto",
    force=False,
):
    """Convert one v1/v2 report or generate a temporary v2 from BRD/MCM/SIP."""
    input_path = Path(input_path).resolve()
    if input_path.suffix.casefold() in DESIGN_SUFFIXES:
        output_path = Path(
            output_path
            or default_design_output(PROJECT_ROOT, input_path, "netlist")
        ).resolve()
        with temporary_v2_report(
            input_path,
            "netlist",
            PROJECT_ROOT,
        ) as report_path:
            result = convert(report_path, output_path, groups, force)
        print("Removed temporary v2 report.")
        return result
    return convert(input_path, output_path, groups, force)


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "input",
        nargs="?",
        type=Path,
        default=PROJECT_ROOT,
        help="one _v1/_v2 report or BRD; default: active PCB folder",
    )
    parser.add_argument("-o", "--output", type=Path)
    parser.add_argument(
        "--groups",
        "--lanes",
        dest="groups",
        default="auto",
        help="PCIe lanes: auto, 0,2, or 0-15",
    )
    parser.add_argument("--force", action="store_true")
    return parser.parse_args()


def main():
    args = parse_args()
    input_path = args.input.resolve()
    if input_path.suffix.casefold() in {
        ".htm",
        ".html",
        ".brd",
        ".mcm",
        ".sip",
    }:
        try:
            convert_source(
                input_path,
                args.output,
                args.groups,
                args.force,
            )
        except NoMatchingSignals as error:
            print("[SKIPPED] {0}: {1}".format(input_path.name, error))
            return NO_SIGNALS
        except Exception as error:
            print("[FAILED] {0}: {1}".format(input_path.name, error))
            return 1
        return 0

    input_path.mkdir(parents=True, exist_ok=True)
    reports = sorted(
        (
            path
            for path in input_path.iterdir()
            if path.is_file() and path.suffix.casefold() in {".htm", ".html"}
        ),
        key=lambda path: path.name.casefold(),
    )
    if not reports:
        print("No netlist .htm/.html files found in: {0}".format(input_path))
        print("Put _v1.htm or _v2.htm files there, then run script/01_Netlist.sh again.")
        return 0

    failures = 0
    skipped = 0
    for report in reports:
        print("[NETLIST] {0}".format(report.name))
        try:
            convert_source(report, groups=args.groups, force=args.force)
        except NoMatchingSignals as error:
            skipped += 1
            print("[SKIPPED] {0}: {1}".format(report.name, error))
        except Exception as error:
            failures += 1
            print("[FAILED] {0}: {1}".format(report.name, error))
    print(
        "Netlist reports: {0}  Succeeded: {1}  Failed: {2}".format(
            len(reports), len(reports) - failures - skipped, failures
        )
    )
    print("Skipped: {0}".format(skipped))
    return 1 if failures else NO_SIGNALS if skipped == len(reports) else 0


if __name__ == "__main__":
    raise SystemExit(extractor_exit(main))
