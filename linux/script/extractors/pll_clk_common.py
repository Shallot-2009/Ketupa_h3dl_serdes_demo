#!/usr/bin/env python3
"""Shared PLL/CLK Cadence report preprocessing for PCB, PKG, and Merge.

The module deliberately keeps the original layout net names in column A so
HFSS cutout can find the physical nets.  Only complete, case-insensitive
``_P/_N`` or ``_DP/_DN`` clock pairs are emitted.  Package fan-out is kept in
full: one BGA/solder pin is written to PIN 1 and every matching DIE bump is
written to PIN 2.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import os
import re
import subprocess
import tempfile
from collections import Counter, defaultdict
from contextlib import contextmanager
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from artifact_paths import artifact_path, NO_SIGNALS, NoMatchingSignals
from cadence_runtime import resolve_cadence_report as _resolve_cadence_report, run_cadence_export


csv.field_size_limit(64 * 1024 * 1024)

FORMAT_VERSION = "ketupa-pll-clk-preprocess-v1"
DESIGN_SUFFIXES = {".brd", ".mcm", ".sip"}
REPORT_SUFFIXES = {".htm", ".html"}
REPORT_CODES = {"netlist": "net", "component": "cmp"}
REPORT_PREFIXES = {
    "netlist": "PLLCLKNetListReport",
    "component": "ComponentReport",
}
PIN_PATTERN = re.compile(r"^(?P<refdes>[^.\s]+)\.(?P<pin>\S+)$")
CLOCK_TOKEN = re.compile(
    r"(?:^|[_\-.])(?:PLL|CLK|CLOCK)[A-Z0-9]*(?=$|[_\-.])", re.IGNORECASE
)
DIFF_SUFFIXES = (
    re.compile(r"(?P<sep>[_\-.]+)D?(?P<polarity>P|N)$", re.IGNORECASE),
    re.compile(
        r"(?P<sep>[_\-.]+)D?(?P<polarity>P|N)\s*[<\[(]"
        r"(?P<index>\d+)[>\])]$",
        re.IGNORECASE,
    ),
)
HEADERS = (
    "Net Name",
    "PIN 1",
    "PIN 2",
    "Signal Type",
    "M Group",
    "Chip Partition",
)
PLACEMENT_HEADERS = (
    "RefDes",
    "Role",
    "Side",
    "Mirror",
    "X",
    "Y",
    "Rotation",
    "Package",
)
HEADER_COLOR = "17365D"
HEADER_ACCENT = "4472C4"
TEXT_COLOR = "243447"
ROW_COLORS = ("D9EAF7", "E2F0D9", "FCE4D6", "E4DFEC", "FFF2CC")


class TableParser(HTMLParser):
    """Collect every HTML table row without depending on browser libraries."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.row = None
        self.cell = None
        self.rows = []

    def handle_starttag(self, tag, attrs):
        del attrs
        tag = tag.casefold()
        if tag == "tr":
            self.row = []
        elif tag in {"td", "th"} and self.row is not None:
            self.cell = []

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)

    def handle_endtag(self, tag):
        tag = tag.casefold()
        if tag in {"td", "th"} and self.cell is not None:
            self.row.append(_clean("".join(self.cell)))
            self.cell = None
        elif tag == "tr" and self.row is not None:
            if self.row:
                self.rows.append(self.row)
            self.row = None


@dataclass(frozen=True)
class ClockNet:
    name: str
    pins: tuple[str, ...]
    group: str
    polarity: str
    pin1: tuple[str, ...]
    pin2: tuple[str, ...]
    partition: str

    def values(self):
        return (
            self.name,
            " ".join(self.pin1),
            " ".join(self.pin2),
            "DIFF_" + self.polarity,
            self.group,
            self.partition,
        )


@dataclass(frozen=True)
class Placement:
    refdes: str
    role: str
    side: str
    mirror: str
    x: str
    y: str
    rotation: str
    package: str

    def values(self):
        return (
            self.refdes,
            self.role,
            self.side,
            self.mirror,
            self.x,
            self.y,
            self.rotation,
            self.package,
        )


def _clean(value):
    return " ".join(str(value).replace("\xa0", " ").split())


def natural_key(value):
    return tuple(
        (0, int(part)) if part.isdigit() else (1, part.casefold())
        for part in re.split(r"(\d+)", str(value))
        if part
    )


def _decode(path):
    raw = Path(path).read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("Cannot decode Cadence report: {0}".format(path))


def _records_from_rows(rows, required_headers):
    requested = tuple(str(value).strip() for value in required_headers)
    required = {value.casefold() for value in requested}
    for index, row in enumerate(rows):
        normalized = [_clean(value).casefold() for value in row]
        if not required.issubset(normalized):
            continue
        positions = {
            header: normalized.index(header.casefold()) for header in requested
        }
        last_column = max(positions.values())
        records = []
        for row_values in rows[index + 1 :]:
            cleaned = [_clean(value) for value in row_values]
            if len(cleaned) <= last_column or not any(cleaned):
                continue
            records.append(
                {header: cleaned[column] for header, column in positions.items()}
            )
        if records:
            return records
    raise ValueError(
        "Cadence report does not contain required columns: {0}".format(
            ", ".join(requested)
        )
    )


def read_report(path, required_headers):
    path = Path(path).resolve()
    if not path.is_file() or path.suffix.casefold() not in REPORT_SUFFIXES:
        raise ValueError("Report must be an existing .htm/.html file: {0}".format(path))
    text = _decode(path)
    if "<table" in text.casefold():
        parser = TableParser()
        parser.feed(text)
        rows = parser.rows
    else:
        rows = list(csv.reader(text.splitlines()))
    return _records_from_rows(rows, required_headers)


def resolve_cadence_report():
    return _resolve_cadence_report()


@contextmanager
def cadence_report(design_path, report_kind):
    design_path = Path(design_path).resolve()
    if not design_path.is_file() or design_path.suffix.casefold() not in DESIGN_SUFFIXES:
        raise ValueError("Invalid Cadence layout file: {0}".format(design_path))
    with tempfile.TemporaryDirectory(prefix="ketupa_pll_clk_") as directory:
        report_path = Path(directory) / "{0}_{1}_v2.htm".format(
            REPORT_PREFIXES[report_kind], design_path.stem
        )
        run_cadence_export(REPORT_CODES[report_kind], design_path, report_path)
        yield report_path


def project_directories(project_root, topology):
    project_root = Path(project_root).resolve()
    mapping = {
        "pcb": ("", "netlist", "placement"),
        "pkg": ("", "netlist", None),
        "merge-pcb": ("", "netlist", "placement"),
        "merge-pkg": ("", "netlist", None),
    }
    if topology not in mapping:
        raise ValueError("Unsupported PLL/CLK topology: {0}".format(topology))
    layout, netlist, placement = mapping[topology]
    return (
        project_root / layout if layout else project_root,
        project_root / netlist,
        None if placement is None else project_root / placement,
    )


def _parse_pin(value):
    match = PIN_PATTERN.fullmatch(str(value).strip())
    if match is None:
        raise ValueError("Invalid RefDes.Pin token: {0}".format(value))
    return match.group("refdes"), match.group("pin")


def _split_pins(value):
    return tuple(
        token for token in re.split(r"[\s,;|]+", str(value).strip()) if token
    )


def clock_identity(net_name):
    """Return a stable group and polarity, or ``None`` for a non-clock net."""
    clean = str(net_name).strip()
    for pattern in DIFF_SUFFIXES:
        match = pattern.search(clean)
        if match is None:
            continue
        base = clean[: match.start()]
        if CLOCK_TOKEN.search(base) is None:
            return None
        group = re.sub(r"[^A-Z0-9]+", "_", base.upper()).strip("_")
        index = match.groupdict().get("index")
        if index:
            group = "{0}_{1}".format(group, index)
        return group, match.group("polarity").upper()
    return None


def _component_pins(pins):
    result = defaultdict(list)
    for pin in pins:
        refdes, _ = _parse_pin(pin)
        result[refdes].append(pin)
    return result


def _pcb_source_component(components):
    def rank(name):
        upper = name.upper()
        if re.fullmatch(r"U\d+(?:[_\-.][A-Z0-9]+)*", upper):
            family = 0
        elif re.fullmatch(r"J\d+(?:[_\-.][A-Z0-9]+)*", upper):
            family = 2
        else:
            family = 1
        return family, natural_key(name)

    return min(components, key=rank)


def _pkg_source_component(component_pins):
    """Select the BGA/solder component; DIE C4 bump pin names are numeric."""
    scored = []
    for component, pins in component_pins.items():
        non_numeric = sum(not _parse_pin(pin)[1].isdigit() for pin in pins)
        scored.append((-non_numeric, -len(pins), natural_key(component), component))
    return min(scored)[-1]


def _build_clock_nets(rows, topology, groups="auto"):
    candidates = []
    seen_names = set()
    pair_members = defaultdict(set)
    for row in rows:
        name = row["Net Name"].strip()
        identity = clock_identity(name)
        if identity is None:
            continue
        key = name.casefold()
        if key in seen_names:
            raise ValueError("Duplicate PLL/CLK network: {0}".format(name))
        seen_names.add(key)
        pins = _split_pins(row["Net Pins"])
        if len(pins) < 2:
            continue
        for pin in pins:
            _parse_pin(pin)
        group, polarity = identity
        candidates.append((name, pins, group, polarity))
        pair_members[group].add(polarity)

    complete_groups = {
        group for group, polarities in pair_members.items() if polarities == {"P", "N"}
    }
    incomplete = sorted(set(pair_members) - complete_groups, key=natural_key)
    if incomplete:
        print(
            "Ignored incomplete PLL/CLK candidates (P/N pair required): {0}".format(
                ", ".join(incomplete)
            )
        )

    selected = str(groups or "auto").strip()
    selected_tokens = [] if selected.casefold() in {"", "auto", "all"} else [
        token.strip().casefold() for token in selected.split(",") if token.strip()
    ]
    records = []
    for name, pins, group, polarity in candidates:
        if group not in complete_groups:
            continue
        if selected_tokens and not any(
            token in group.casefold() or token in name.casefold()
            for token in selected_tokens
        ):
            continue
        by_component = _component_pins(pins)
        if len(by_component) < 2:
            raise ValueError(
                "PLL/CLK net must span at least two components: {0}".format(name)
            )
        if topology.endswith("pkg"):
            source_component = _pkg_source_component(by_component)
        else:
            source_component = _pcb_source_component(by_component)
        pin1 = tuple(by_component[source_component])
        pin2 = tuple(
            pin
            for component in sorted(by_component, key=natural_key)
            if component != source_component
            for pin in by_component[component]
        )
        if not pin1 or not pin2:
            raise ValueError("Cannot split PLL/CLK endpoints for: {0}".format(name))
        records.append(
            ClockNet(
                name=name,
                pins=pins,
                group=group,
                polarity=polarity,
                pin1=pin1,
                pin2=pin2,
                partition="{0}_{1}".format(
                    "PKG" if topology.endswith("pkg") else "PCB",
                    source_component.upper(),
                ),
            )
        )
    if not records:
        if not seen_names:
            raise NoMatchingSignals("No differential PLL/CLK signals were found")
        raise ValueError(
            "No complete differential PLL/CLK _P/_N or _DP/_DN pairs were found"
        )

    counts = Counter((record.group, record.polarity) for record in records)
    invalid = [
        "{0}:{1}={2}".format(group, polarity, count)
        for (group, polarity), count in counts.items()
        if count != 1
    ]
    if invalid:
        raise ValueError(
            "Each PLL/CLK group must contain one physical P net and one N net: {0}".format(
                "; ".join(invalid)
            )
        )
    return sorted(
        records,
        key=lambda record: (
            natural_key(record.group),
            0 if record.polarity == "P" else 1,
            natural_key(record.name),
        ),
    )


def _style_sheet(sheet, row_count, widths):
    for cell in sheet[1]:
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_COLOR)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=Side(style="medium", color=HEADER_ACCENT))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:{0}{1}".format(
        chr(ord("A") + len(widths) - 1), row_count + 1
    )
    sheet.sheet_view.showGridLines = False
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(ord("A") + column - 1)].width = width
    for row in range(2, row_count + 2):
        color = ROW_COLORS[((row - 2) // 2) % len(ROW_COLORS)]
        for column in range(1, len(widths) + 1):
            cell = sheet.cell(row, column)
            cell.font = Font(name="Microsoft YaHei", size=9, color=TEXT_COLOR)
            cell.fill = PatternFill("solid", fgColor=color)
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border = Border(bottom=Side(style="hair", color="D6DEE8"))


def _fingerprint(values):
    digest = hashlib.sha256(FORMAT_VERSION.encode("ascii"))
    for row in values:
        for value in row:
            digest.update(str(value).encode("utf-8"))
            digest.update(b"\0")
    return digest.hexdigest().upper()


def _atomic_save(workbook, output_path, expected_rows, expected_headers):
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_name(".{0}.tmp.xlsx".format(output_path.stem))
    try:
        workbook.save(temporary)
        check = load_workbook(temporary, read_only=True, data_only=True)
        values = list(check[check.sheetnames[0]].iter_rows(values_only=True))
        check.close()
        if not values or tuple(values[0]) != tuple(expected_headers):
            raise AssertionError("Generated workbook header validation failed")
        normalized = [tuple("" if v is None else v for v in row) for row in values[1:]]
        if normalized != list(expected_rows):
            raise AssertionError("Generated workbook row validation failed")
        temporary.replace(output_path)
    finally:
        workbook.close()
        if temporary.exists():
            temporary.unlink()


def _source_stem(source):
    stem = re.sub(r"_v[12]$", "", Path(source).stem, flags=re.IGNORECASE)
    for prefix in ("PLLCLKNetListReport_", "ComponentReport_"):
        if stem.casefold().startswith(prefix.casefold()):
            return stem[len(prefix) :]
    return stem


def _legacy_convert_netlist(source, project_root, topology, output=None, groups="auto", force=False):
    source = Path(source).resolve()
    _, netlist_dir, _ = project_directories(project_root, topology)
    netlist_dir.mkdir(parents=True, exist_ok=True)
    output_path = (
        Path(output).resolve()
        if output
        else artifact_path(
            project_root, Path(_source_stem(source)), "netlist", "pll_clk"
        )
    )
    if output_path.suffix.casefold() != ".xlsx":
        raise ValueError("Netlist output must use .xlsx: {0}".format(output_path))

    def build(report_path):
        rows = read_report(report_path, ("Net Name", "Net Pins"))
        return _build_clock_nets(rows, topology, groups)

    if source.suffix.casefold() in DESIGN_SUFFIXES:
        with cadence_report(source, "netlist") as report_path:
            records = build(report_path)
    elif source.suffix.casefold() in REPORT_SUFFIXES:
        records = build(source)
    else:
        raise ValueError("Netlist input must be BRD/MCM/SIP/HTM/HTML: {0}".format(source))
    values = [record.values() for record in records]
    fingerprint = _fingerprint(values)
    if not force and output_path.is_file():
        current = load_workbook(output_path, read_only=True)
        unchanged = current.properties.keywords == "source-sha256=" + fingerprint
        current.close()
        if unchanged:
            print("Unchanged: {0}".format(output_path))
            return output_path
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ALL"
    sheet.append(HEADERS)
    for row in values:
        sheet.append(row)
    _style_sheet(sheet, len(values), (38, 30, 58, 14, 24, 24))
    workbook.properties.title = "Ketupa PLL/CLK Differential Netlist"
    workbook.properties.keywords = "source-sha256=" + fingerprint
    _atomic_save(workbook, output_path, values, HEADERS)
    fanout = max(len(record.pin2) for record in records)
    print("Created: {0}".format(output_path))
    print(
        "PLL/CLK differential nets: {0}; pairs: {1}; maximum fan-out: {2}".format(
            len(records), len({record.group for record in records}), fanout
        )
    )
    return output_path


def convert_netlist(source, project_root, topology, output=None, groups="auto", force=False):
    """Use legacy clock extraction first, then the approved universal profile."""
    source = Path(source).resolve()
    output_path = (
        Path(output).resolve()
        if output
        else artifact_path(
            project_root, Path(_source_stem(source)), "netlist", "pll_clk"
        )
    )
    from script.data_adapter.runtime import active_settings, normalize_report

    settings = active_settings()

    def adapter():
        selected_topology = "pkg" if str(topology).endswith("pkg") else "pcb"
        if source.suffix.casefold() in DESIGN_SUFFIXES:
            with cadence_report(source, "netlist") as report_path:
                return normalize_report(
                    report_path,
                    output_path,
                    family="pll_clk",
                    topology=selected_topology,
                    groups=groups,
                    force=force,
                )
        return normalize_report(
            source,
            output_path,
            family="pll_clk",
            topology=selected_topology,
            groups=groups,
            force=force,
        )

    if settings["mode"] == "universal":
        return adapter()
    try:
        return _legacy_convert_netlist(
            source, project_root, topology, output_path, groups, force
        )
    except (OSError, ValueError, RuntimeError) as legacy_error:
        if settings["mode"] == "legacy":
            raise
        try:
            print("Legacy extractor did not accept this naming: {0}".format(legacy_error))
            print("Trying the approved universal profile...", flush=True)
            return adapter()
        except (OSError, ValueError, RuntimeError) as adapter_error:
            raise ValueError(
                "Both extraction paths rejected the input. Legacy: {0}. "
                "Universal: {1}".format(legacy_error, adapter_error)
            ) from adapter_error


def _netlist_roles(netlist_path):
    workbook = load_workbook(netlist_path, read_only=True, data_only=True)
    sheet = next(
        (value for value in workbook.worksheets if value.title.casefold() == "all"),
        workbook.worksheets[0],
    )
    roles = {}
    priority = {"CLOCK_RECEIVER": 0, "CLOCK_DRIVER": 1, "CHIP": 2, "CONNECTOR": 3}
    try:
        for row in sheet.iter_rows(min_row=2, min_col=2, max_col=3, values_only=True):
            for column, cell in enumerate(row, start=1):
                for pin in _split_pins(cell or ""):
                    refdes, _ = _parse_pin(pin)
                    upper = refdes.upper()
                    if re.fullmatch(r"J\d+(?:[_\-.][A-Z0-9]+)*", upper):
                        role = "CONNECTOR"
                    elif column == 1 and upper.startswith("U"):
                        role = "CHIP"
                    elif column == 1:
                        role = "CLOCK_DRIVER"
                    else:
                        role = "CLOCK_RECEIVER"
                    previous = roles.get(upper)
                    if previous is None or priority[role] > priority[previous]:
                        roles[upper] = role
    finally:
        workbook.close()
    if not roles:
        raise ValueError("Netlist workbook contains no PLL/CLK endpoint pins")
    return roles


def convert_placement(source, netlist_path, project_root, topology, output=None, force=False):
    source = Path(source).resolve()
    _, _, placement_dir = project_directories(project_root, topology)
    if placement_dir is None:
        raise ValueError("Package-only PLL/CLK preprocessing does not create Placement")
    placement_dir.mkdir(parents=True, exist_ok=True)
    netlist_path = Path(netlist_path).resolve()
    if not netlist_path.is_file():
        raise FileNotFoundError("PLL/CLK netlist workbook not found: {0}".format(netlist_path))
    output_path = (
        Path(output).resolve()
        if output
        else artifact_path(
            project_root, Path(_source_stem(source)), "placement", "pll_clk"
        )
    )
    required = ("REFDES", "SYM_MIRROR", "SYM_X", "SYM_Y", "SYM_ROTATE", "COMP_PACKAGE")

    def build(report_path):
        roles = _netlist_roles(netlist_path)
        records = []
        seen = set()
        for row in read_report(report_path, required):
            refdes = row["REFDES"].strip().upper()
            if refdes not in roles:
                continue
            seen.add(refdes)
            mirror = row["SYM_MIRROR"].strip().upper()
            if mirror == "NO":
                side = "TOP"
            elif mirror == "YES":
                side = "BOTTOM"
            else:
                raise ValueError("Unsupported SYM_MIRROR for {0}: {1}".format(refdes, mirror))
            if not any(row.get(field, "").strip() for field in required[1:]):
                continue
            records.append(
                Placement(
                    refdes,
                    roles[refdes],
                    side,
                    mirror,
                    row["SYM_X"],
                    row["SYM_Y"],
                    row["SYM_ROTATE"],
                    row["COMP_PACKAGE"],
                )
            )
        missing = sorted(set(roles) - seen, key=natural_key)
        if missing:
            raise ValueError(
                "PLL/CLK endpoint RefDes missing from component report: {0}".format(
                    ", ".join(missing)
                )
            )
        if not records:
            raise ValueError("No placed PLL/CLK endpoint components were found")
        return sorted(records, key=lambda record: natural_key(record.refdes))

    if source.suffix.casefold() in DESIGN_SUFFIXES:
        with cadence_report(source, "component") as report_path:
            records = build(report_path)
    elif source.suffix.casefold() in REPORT_SUFFIXES:
        records = build(source)
    else:
        raise ValueError("Placement input must be BRD/MCM/SIP/HTM/HTML: {0}".format(source))
    values = [record.values() for record in records]
    fingerprint = _fingerprint(values)
    if not force and output_path.is_file():
        current = load_workbook(output_path, read_only=True)
        unchanged = current.properties.keywords == "source-sha256=" + fingerprint
        current.close()
        if unchanged:
            print("Unchanged: {0}".format(output_path))
            return output_path
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Placement"
    sheet.append(PLACEMENT_HEADERS)
    for row in values:
        sheet.append(row)
    _style_sheet(sheet, len(values), (18, 24, 12, 12, 16, 16, 14, 28))
    workbook.properties.title = "Ketupa PLL/CLK PCB Placement"
    workbook.properties.keywords = "source-sha256=" + fingerprint
    _atomic_save(workbook, output_path, values, PLACEMENT_HEADERS)
    counts = Counter(record.role for record in records)
    print("Netlist: {0}".format(netlist_path))
    print("Created: {0}".format(output_path))
    print(
        "Placement rows: {0}; {1}".format(
            len(records),
            ", ".join("{0}={1}".format(role, counts[role]) for role in sorted(counts)),
        )
    )
    return output_path


def collect_sources(project_root, topology, selected):
    layout_dir, _, _ = project_directories(project_root, topology)
    layout_dir.mkdir(parents=True, exist_ok=True)
    sources = [Path(value).resolve() for value in selected]
    if not sources:
        sources = sorted(
            (
                path.resolve()
                for path in layout_dir.iterdir()
                if path.is_file() and path.suffix.casefold() in DESIGN_SUFFIXES
            ),
            key=lambda path: path.name.casefold(),
        )
    for source in sources:
        if not source.is_file() or source.suffix.casefold() not in DESIGN_SUFFIXES | REPORT_SUFFIXES:
            raise ValueError("Unsupported or missing PLL/CLK input: {0}".format(source))
    return sources


def netlist_cli(project_root, topology, argv=None):
    parser = argparse.ArgumentParser(description="Generate PLL/CLK differential netlist Excel")
    parser.add_argument("sources", nargs="*")
    parser.add_argument("--output")
    parser.add_argument("--groups", default="auto")
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args(argv)
    sources = collect_sources(project_root, topology, args.sources)
    if not sources:
        print("No layout files found in: {0}".format(project_directories(project_root, topology)[0]))
        return 0
    if args.output and len(sources) != 1:
        parser.error("--output requires exactly one input")
    failures = []
    skipped = 0
    for source in sources:
        try:
            convert_netlist(source, project_root, topology, args.output, args.groups, args.force)
        except NoMatchingSignals as error:
            skipped += 1
            print("[SKIPPED] {0}: {1}".format(source.name, error))
        except Exception as error:
            failures.append((source, error))
            print("[FAILED] {0}: {1}".format(source.name, error))
    print("Sources: {0}  Succeeded: {1}  Skipped: {2}  Failed: {3}".format(len(sources), len(sources) - len(failures) - skipped, skipped, len(failures)))
    return 1 if failures else NO_SIGNALS if skipped == len(sources) else 0


def placement_cli(project_root, topology, argv=None):
    parser = argparse.ArgumentParser(description="Generate PLL/CLK endpoint Placement Excel")
    parser.add_argument("source", nargs="?")
    parser.add_argument("--netlist")
    parser.add_argument("--output")
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args(argv)
    sources = collect_sources(project_root, topology, [args.source] if args.source else [])
    if len(sources) != 1:
        raise RuntimeError("Placement requires exactly one layout/report input; found {0}".format(len(sources)))
    source = sources[0]
    _, netlist_dir, _ = project_directories(project_root, topology)
    netlist_path = (
        Path(args.netlist).resolve()
        if args.netlist
        else artifact_path(
            project_root, Path(_source_stem(source)), "netlist", "pll_clk"
        )
    )
    convert_placement(source, netlist_path, project_root, topology, args.output, args.force)
    return 0


def pipeline_cli(project_root, topology, argv=None):
    parser = argparse.ArgumentParser(description="Run PLL/CLK netlist and placement preprocessing")
    parser.add_argument("sources", nargs="*")
    parser.add_argument("--groups", default="auto")
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args(argv)
    sources = collect_sources(project_root, topology, args.sources)
    if not sources:
        print("No layout files found in: {0}".format(project_directories(project_root, topology)[0]))
        return 0
    failures = []
    for source in sources:
        try:
            print("\n[1/2 NETLIST] {0}".format(source.name))
            netlist = convert_netlist(source, project_root, topology, groups=args.groups, force=args.force)
            print("[2/2 PLACEMENT] {0}".format(source.name))
            convert_placement(source, netlist, project_root, topology, force=args.force)
            print("[OK] {0}".format(source.name))
        except Exception as error:
            failures.append((source, error))
            print("[FAILED] {0}: {1}".format(source.name, error))
    print("\nDesigns: {0}  Succeeded: {1}  Failed: {2}".format(len(sources), len(sources) - len(failures), len(failures)))
    return 1 if failures else 0
