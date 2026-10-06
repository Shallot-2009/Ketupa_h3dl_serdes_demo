#!/usr/bin/env python3
"""Convert SerDes reports or MCM/SIP designs to DIE-first grouped Excel."""

from __future__ import annotations

import argparse
import csv
import hashlib
import os
import re
import subprocess
import tempfile
from collections import Counter
from contextlib import contextmanager
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from artifact_paths import artifact_path, NO_SIGNALS, NoMatchingSignals, extractor_exit
from cadence_runtime import resolve_cadence_export_command as _resolve_cadence_export_command, run_cadence_export
from signal_identity import parse_serdes_network_name

csv.field_size_limit(64 * 1024 * 1024)

FORMAT_VERSION = "ketupa-serdes-netlist-v4"
SORT_CONTRACT = "DIE -> SDS -> RX/TX lane -> P/N (P then N)"
PROJECT_ROOT = Path(os.environ["KETUPA_WORKFLOW_INPUT"]).resolve()
DESIGN_SUFFIXES = {".aedb", ".mcm", ".sip"}
REPORT_CODES = {"netlist": "net", "component": "cmp"}
REPORT_PREFIXES = {
    "netlist": "SerDesNetListReport",
    "component": "ComponentReport",
}
VARIANT_SUFFIX = re.compile(r"_v(?P<version>[12])$", re.I)
DIRECTION_RANK = {"RX": 0, "TX": 1}
POLARITY_RANK = {"P": 0, "N": 1}
LEGACY_CHIP_SUFFIX = re.compile(r"(?:_|-|\.)G(?P<index>\d*)$", re.I)
PIN = re.compile(r"^(?P<refdes>[^.\s]+)\.(?P<pin>\S+)$")
HEADERS = ("Net Name", "PIN 1", "PIN 2", "Signal Type", "M Group", "Chip Partition")

HEADER_COLOR = "17365D"
HEADER_ACCENT = "4472C4"
TEXT_COLOR = "243447"
GRID_COLOR = "D6DEE8"
BLOCK_COLORS = (
    "D9EAF7",
    "E2F0D9",
    "FCE4D6",
    "E4DFEC",
    "FFF2CC",
    "DDEBF7",
    "EADCF8",
    "F4CCCC",
    "D9EAD3",
    "D0E0E3",
    "FCE5CD",
    "DDEBF1",
)
DIE_COLORS = (
    "E2F0D9",
    "FCE4D6",
    "DDEBF7",
    "E4DFEC",
    "FFF2CC",
    "D9EAD3",
)
CHANNEL_ACCENTS = ("4472C4", "70AD47", "ED7D31", "A64D79", "5B9BD5")


class TableParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.row = None
        self.cell = None
        self.rows = []

    def handle_starttag(self, tag, attrs):
        del attrs
        tag = tag.casefold()
        if tag == "tr":
            self.row = []
        elif tag in {"td", "th"} and self.row is not None:
            self.cell = []

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)

    def handle_endtag(self, tag):
        tag = tag.casefold()
        if tag in {"td", "th"} and self.cell is not None:
            self.row.append(" ".join("".join(self.cell).replace("\xa0", " ").split()))
            self.cell = None
        elif tag == "tr" and self.row is not None:
            if self.row:
                self.rows.append(self.row)
            self.row = None


def _clean_cell(value):
    return " ".join(str(value).replace("\xa0", " ").split())


def _decode_report(path):
    raw = Path(path).read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("Cannot decode Cadence report: {0}".format(path))


def report_variant(path, text=None):
    path = Path(path)
    text = _decode_report(path) if text is None else text
    content_variant = "v1" if "<table" in text.casefold() else "v2"
    suffix_match = VARIANT_SUFFIX.search(path.stem)
    if suffix_match is not None:
        named_variant = "v" + suffix_match.group("version")
        if named_variant != content_variant:
            raise ValueError(
                "Report suffix/content mismatch for {0}: name={1}, content={2}".format(
                    path, named_variant, content_variant
                )
            )
    return content_variant


def canonical_report_stem(path):
    return VARIANT_SUFFIX.sub("", Path(path).stem)


def canonical_xlsx_path(path, output_path=None):
    path = Path(path).resolve()
    selected = (
        Path(output_path).resolve()
        if output_path is not None
        else path.with_name(canonical_report_stem(path) + ".xlsx")
    )
    if selected.suffix.casefold() != ".xlsx":
        raise ValueError("Output must use .xlsx: {0}".format(selected))
    return selected


def _records_from_rows(rows, required_headers):
    requested = tuple(str(header).strip() for header in required_headers)
    required = {header.casefold() for header in requested}
    for index, row in enumerate(rows):
        normalized = [_clean_cell(value).casefold() for value in row]
        if not required.issubset(set(normalized)):
            continue
        positions = {
            header: normalized.index(header.casefold()) for header in requested
        }
        last_column = max(positions.values())
        records = []
        for values in rows[index + 1 :]:
            cleaned = [_clean_cell(value) for value in values]
            if len(cleaned) <= last_column or not any(cleaned):
                continue
            records.append(
                {header: cleaned[column] for header, column in positions.items()}
            )
        if records:
            return records
    raise ValueError(
        "Cadence report does not contain the required columns: {0}".format(
            ", ".join(requested)
        )
    )


def read_report_table(path, required_headers):
    path = Path(path).resolve()
    if path.suffix.casefold() not in {".htm", ".html"} or not path.is_file():
        raise ValueError("Report must be an existing .htm/.html file: {0}".format(path))
    text = _decode_report(path)
    if report_variant(path, text) == "v1":
        parser = TableParser()
        parser.feed(text)
        rows = parser.rows
    else:
        rows = list(csv.reader(text.splitlines()))
    return _records_from_rows(rows, required_headers)


def resolve_cadence_export_command():
    return _resolve_cadence_export_command()


def default_design_output(project_root, design_path, report_kind):
    return artifact_path(project_root, design_path, report_kind, "serdes")


@contextmanager
def temporary_v2_report(design_path, report_kind, project_root=None):
    del project_root
    design_path = Path(design_path).resolve()
    if design_path.suffix.casefold() not in DESIGN_SUFFIXES or not design_path.is_file():
        raise ValueError("Design must be an existing .mcm/.sip file: {0}".format(design_path))
    with tempfile.TemporaryDirectory(prefix="ketupa_cadence_report_") as directory:
        report_path = Path(directory) / "{0}_{1}_v2.htm".format(
            REPORT_PREFIXES[report_kind], design_path.stem
        )
        run_cadence_export(REPORT_CODES[report_kind], design_path, report_path)
        report_variant(report_path)
        try:
            yield report_path
        finally:
            pass


def read_table(path, required_headers):
    return read_report_table(path, required_headers)


def _rgb(value):
    value = value.lstrip("#")
    return tuple(int(value[index : index + 2], 16) for index in (0, 2, 4))


def _hex(values):
    return "".join("{0:02X}".format(max(0, min(255, value))) for value in values)


def _darken(color, amount=0.28):
    return _hex(tuple(round(value * (1 - amount)) for value in _rgb(color)))


def _channel_color(channel):
    number = sum(ord(character) for character in str(channel).upper())
    return CHANNEL_ACCENTS[number % len(CHANNEL_ACCENTS)]


def color_block_key(record):
    net = record.net
    return net.sds, net.die, net.direction, net.lane


def style_netlist_sheet(sheet, records):
    for cell in sheet[1][:6]:
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_COLOR)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=Side(style="medium", color=HEADER_ACCENT))
    sheet.row_dimensions[1].height = 28
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:F{0}".format(sheet.max_row)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90
    for column, width in zip("ABCDEF", (34, 30, 44, 14, 12, 18)):
        sheet.column_dimensions[column].width = width
    if not records:
        sheet.sheet_properties.tabColor = HEADER_COLOR
        return

    block_keys = list(dict.fromkeys(color_block_key(record) for record in records))
    block_colors = {
        key: BLOCK_COLORS[index % len(BLOCK_COLORS)]
        for index, key in enumerate(block_keys)
    }
    previous = None
    for row_index, record in enumerate(records, start=2):
        net = record.net
        block_color = block_colors[color_block_key(record)]
        die_color = DIE_COLORS[(net.die - 1) % len(DIE_COLORS)]
        channel_color = _channel_color("SDS{0}".format(net.sds))
        sds_changed = previous is None or previous.net.sds != net.sds
        die_changed = sds_changed or previous.net.die != net.die
        group_changed = (
            die_changed
            or previous.net.direction != net.direction
            or previous.net.lane != net.lane
        )
        block_changed = group_changed or previous.net.direction != net.direction
        next_record = records[row_index - 1] if row_index - 1 < len(records) else None
        block_ends = next_record is None or color_block_key(
            next_record
        ) != color_block_key(record)
        top_color = _darken(die_color) if die_changed else _darken(block_color)
        for column_index in range(1, 7):
            cell = sheet.cell(row_index, column_index)
            cell.font = Font(
                name="Microsoft YaHei",
                size=9,
                color=TEXT_COLOR,
                bold=column_index in {1, 5, 6},
            )
            cell.fill = PatternFill(
                "solid", fgColor=die_color if column_index == 6 else block_color
            )
            cell.alignment = Alignment(
                horizontal="center" if column_index >= 4 else "left",
                vertical="center",
            )
            cell.border = Border(
                left=(
                    Side(style="medium", color=channel_color)
                    if column_index == 1
                    else Side()
                ),
                top=(
                    Side(
                        style="medium",
                        color=top_color
                        if die_changed or group_changed
                        else channel_color,
                    )
                    if block_changed
                    else Side()
                ),
                bottom=Side(
                    style="medium" if block_ends else "hair",
                    color=_darken(block_color) if block_ends else GRID_COLOR,
                ),
            )
        sheet.row_dimensions[row_index].height = 19
        previous = record
    sheet.sheet_properties.tabColor = (
        HEADER_COLOR
        if sheet.title == "ALL"
        else _darken(BLOCK_COLORS[records[0].net.lane % len(BLOCK_COLORS)], 0.15)
    )


@dataclass(frozen=True)
class NetRecord:
    name: str
    pins: tuple[str, ...]
    sds: int | None
    direction: str
    lane: int
    polarity: str
    signal_type: str
    partition: str
    die: int


@dataclass(frozen=True)
class ExcelRecord:
    net: NetRecord
    pin1: str
    pin2: str

    def values(self):
        group_name = (
            "SDS{0}_{1}{2}".format(
                self.net.sds,
                self.net.direction,
                self.net.lane,
            )
            if self.net.sds is not None
            else "{0}{1}".format(self.net.direction, self.net.lane)
        )
        return (
            self.net.name,
            self.pin1,
            self.pin2,
            self.net.signal_type,
            group_name,
            self.net.partition,
        )


def natural_key(value):
    return tuple(
        (0, int(part)) if part.isdigit() else (1, part.casefold())
        for part in re.split(r"(\d+)", str(value))
        if part
    )


def refdes(pin):
    match = PIN.fullmatch(pin)
    if not match:
        raise ValueError("Invalid RefDes.Pin token: {0}".format(pin))
    return match.group("refdes")


def parse_net_name(name):
    parsed = parse_serdes_network_name(name)
    if parsed is None:
        return None
    sds, direction, lane, polarity, die = parsed
    return (
        sds,
        direction,
        lane,
        polarity,
        "DIFF_{0}".format(polarity),
        "DIE_{0}".format(die),
        die,
    )


def component_die_number(value):
    """Return an explicit component DIE number; shared package parts return None."""
    text = str(value).strip()
    match = re.fullmatch(r"U(?P<index>\d+)(?:[_\-.].*)?", text, re.I)
    if match and 1 <= int(match.group("index")) <= 5:
        return int(match.group("index"))
    match = re.search(
        r"(?:^|[_\-.])D(?:IE)?[_\-.]?(?P<index>[1-5])(?:$|[_\-.])",
        text,
        re.I,
    )
    if match:
        return int(match.group("index"))
    legacy = LEGACY_CHIP_SUFFIX.search(text)
    if legacy:
        return int(legacy.group("index") or "1") + 1
    return None


def parse_group_selection(value):
    text = str(value).strip().casefold()
    if text in {"", "auto", "all"}:
        return None
    selected = set()
    for token in text.split(","):
        token = token.strip().removeprefix("lane")
        if "-" in token:
            start, end = token.split("-", 1)
            start_number = int(start.removeprefix("lane"))
            end_number = int(end.removeprefix("lane"))
            if start_number > end_number:
                raise ValueError("Descending SerDes lane range: {0}".format(token))
            selected.update(range(start_number, end_number + 1))
        else:
            selected.add(int(token))
    invalid = sorted(lane for lane in selected if lane < 0 or lane > 15)
    if invalid:
        raise ValueError("SerDes lanes must be in the range 0..15: {0}".format(invalid))
    return selected


def load_net_records(path, selected_groups=None):
    records = []
    seen = set()
    skipped_single_pin = []
    for row in read_table(path, ("Net Name", "Net Pins")):
        name = row["Net Name"].strip()
        parsed = parse_net_name(name)
        if parsed is None:
            continue
        sds, direction, lane, polarity, category, partition, die = parsed
        if selected_groups is not None and lane not in selected_groups:
            continue
        key = name.casefold()
        if key in seen:
            raise ValueError("Duplicate SerDes network: {0}".format(name))
        seen.add(key)
        pins = tuple(
            token for token in re.split(r"[\s,;|]+", row["Net Pins"].strip()) if token
        )
        if len(pins) < 2:
            skipped_single_pin.append(name)
            continue
        mismatched = []
        for pin in pins:
            component = refdes(pin)
            component_die = component_die_number(component)
            if component_die is not None and component_die != die:
                mismatched.append(component)
        mismatched = sorted(set(mismatched), key=natural_key)
        if mismatched:
            raise ValueError(
                "Network/RefDes chip partition mismatch for {0}: {1}".format(
                    name, ", ".join(mismatched)
                )
            )
        records.append(
            NetRecord(
                name,
                pins,
                sds,
                direction,
                lane,
                polarity,
                category,
                partition,
                die,
            )
        )
    if not records:
        raise NoMatchingSignals("No SerDes SDS TX/RX lane networks were found in {0}".format(path))
    if skipped_single_pin:
        print(
            "Skipped single-pin package nets: {0} ({1})".format(
                len(skipped_single_pin), ", ".join(skipped_single_pin)
            )
        )
    validate_differential_pairs(records)
    return records


def validate_differential_pairs(records):
    """Require exactly one case-insensitive P/N net for every selected lane."""
    pairs = {}
    for record in records:
        key = (record.sds, record.partition, record.direction, record.lane)
        pairs.setdefault(key, []).append(record.polarity)
    invalid = [
        "{0}:{1}:{2}{3}={4}".format(
            "SDS{0}".format(sds) if sds is not None else "LEGACY",
            partition,
            direction,
            lane,
            sorted(polarities),
        )
        for (sds, partition, direction, lane), polarities in sorted(
            pairs.items(),
            key=lambda item: (
                (0, item[0][0]) if item[0][0] is not None else (1, 0),
                natural_key(item[0][1]),
                DIRECTION_RANK[item[0][2]],
                item[0][3],
            ),
        )
        if sorted(polarities) != ["N", "P"]
    ]
    if invalid:
        raise ValueError(
            "Every SerDes lane must contain exactly one P and one N net: {0}".format(
                "; ".join(invalid)
            )
        )


def endpoint_group_key(record):
    """Keep legacy pairs independent from SDS package/DIE partitions."""
    if record.sds is None:
        return (None, record.partition, record.direction, record.lane)
    return (record.sds, record.partition, None, None)


def primary_refdes_by_partition(records):
    partitions = sorted(
        {endpoint_group_key(record) for record in records},
        key=lambda item: (
            (0, item[0]) if item[0] is not None else (1, 0),
            natural_key(item[1]),
            DIRECTION_RANK.get(item[2], 0),
            item[3] if item[3] is not None else -1,
        ),
    )
    coverage = {partition: Counter() for partition in partitions}
    first_seen = {partition: {} for partition in partitions}
    for row_number, record in enumerate(records):
        partition = endpoint_group_key(record)
        for pin_index, component in enumerate(
            dict.fromkeys(refdes(pin) for pin in record.pins)
        ):
            coverage[partition][component] += 1
            first_seen[partition].setdefault(
                component,
                (row_number, pin_index),
            )
    return {
        partition: min(
            counts,
            key=lambda name: (
                -counts[name],
                first_seen[partition][name],
                natural_key(name),
            ),
        )
        for partition, counts in coverage.items()
        if counts
    }


def build_excel_records(records):
    primary = primary_refdes_by_partition(records)
    output = []
    for record in records:
        controller = primary[endpoint_group_key(record)]
        controller_pins = [pin for pin in record.pins if refdes(pin) == controller]
        target_pins = [pin for pin in record.pins if refdes(pin) != controller]
        if not controller_pins or not target_pins:
            raise ValueError(
                "Cannot split BGA and C4 endpoints for {0}".format(record.name)
            )
        output.append(
            ExcelRecord(record, " ".join(controller_pins), " ".join(target_pins))
        )
    return output


def _edb_component_role(refdes, component):
    """Classify only the two physical package endpoints used by this flow."""
    name = str(refdes).strip().upper()
    component_type = str(getattr(component, "type", "")).strip().casefold()
    if component_type == "io" or name.startswith("BGA"):
        return "BGA"
    if component_type == "ic" or name.startswith(("DIE", "C4")):
        return "C4"
    return ""


def _configure_pyedb_installation(version):
    """Expose the selected AEDT root in the form required by PyEDB.

    Product runtime selection historically exported ``ANSYSEM_ROOT`` while
    PyEDB 0.83 discovers installations only through ``ANSYSEM_ROOTxxx``.
    Supplying both keeps the extractor compatible with AEDT 2025 and 2026 and
    also works when this script is started directly instead of through main.
    """
    match = re.fullmatch(r"20(?P<year>\d{2})\.(?P<release>\d)", str(version))
    if match is None:
        raise ValueError("Unsupported PyEDB/AEDT version: {0}".format(version))
    root_value = os.environ.get("ANSYSEM_ROOT", "").strip()
    if not root_value:
        settings_file = Path(__file__).resolve().with_name("cds_env")
        if settings_file.is_file():
            for raw_line in settings_file.read_text(
                encoding="utf-8-sig", errors="replace"
            ).splitlines():
                line = raw_line.strip()
                if line.startswith("KETUPA_ANSYSEDT="):
                    root_value = line.split("=", 1)[1].strip().strip('"')
                    break
    root = Path(os.path.expandvars(root_value)).expanduser().resolve()
    if not root.is_dir() or not (root / "ansysedt").is_file():
        raise FileNotFoundError(
            "Selected AEDT root is unavailable for PyEDB: {0}".format(root)
        )
    variable = "ANSYSEM_ROOT{0}{1}".format(
        match.group("year"), match.group("release")
    )
    os.environ["ANSYSEM_ROOT"] = str(root)
    os.environ[variable] = str(root)
    return variable, root


def load_aedb_excel_records(path, selected_groups=None):
    """Read authentic BGA-to-die pins from one native Ansys EDB."""
    path = Path(path).expanduser().resolve()
    if (
        path.suffix.casefold() != ".aedb"
        or not path.is_dir()
        or not (path / "edb.def").is_file()
    ):
        raise ValueError(
            "AEDB input must be a directory containing edb.def: {0}".format(path)
        )

    from lib.core.cache_policy import (
        _install_version_environment,
        selected_edb_version,
    )

    version = selected_edb_version()
    _install_version_environment(version, os.environ)
    environment_name, environment_root = _configure_pyedb_installation(version)
    print(
        "PyEDB runtime: version={0}; {1}={2}".format(
            version, environment_name, environment_root
        ),
        flush=True,
    )
    from pyedb import Edb

    grpc = tuple(int(value) for value in version.split(".")) >= (2026, 1)
    edb = Edb(
        str(path),
        isreadonly=True,
        version=version,
        grpc=grpc,
        in_memory=False,
    )
    try:
        components = []
        for refdes, component in sorted(
            edb.components.instances.items(), key=lambda item: natural_key(item[0])
        ):
            role = _edb_component_role(refdes, component)
            if role:
                components.append((str(refdes), role, component))
        if not any(role == "BGA" for _, role, _ in components):
            raise ValueError("AEDB contains no BGA/IO package endpoint")
        if not any(role == "C4" for _, role, _ in components):
            raise ValueError("AEDB contains no DIE/IC C4 package endpoint")

        net_records = []
        excel_records = []
        for net_name in sorted(edb.nets.nets, key=natural_key):
            parsed = parse_net_name(str(net_name))
            if parsed is None:
                continue
            sds, direction, lane, polarity, category, partition, die = parsed
            if selected_groups is not None and lane not in selected_groups:
                continue
            endpoint_pins = {"BGA": [], "C4": []}
            endpoint_components = {"BGA": set(), "C4": set()}
            for refdes, role, component in components:
                for pin_name, pin in component.pins.items():
                    if str(getattr(pin, "net_name", "")).casefold() != str(
                        net_name
                    ).casefold():
                        continue
                    endpoint_pins[role].append(
                        "{0}.{1}".format(refdes, str(pin_name).strip())
                    )
                    endpoint_components[role].add(refdes.casefold())
            if not endpoint_pins["BGA"] or not endpoint_pins["C4"]:
                raise ValueError(
                    "SerDes net {0} does not reach both BGA and C4 endpoints".format(
                        net_name
                    )
                )
            for role in ("BGA", "C4"):
                if len(endpoint_components[role]) != 1:
                    raise ValueError(
                        "SerDes net {0} reaches multiple {1} components: {2}".format(
                            net_name,
                            role,
                            sorted(endpoint_components[role]),
                        )
                    )
                endpoint_pins[role].sort(key=natural_key)
            record = NetRecord(
                str(net_name),
                tuple(endpoint_pins["BGA"] + endpoint_pins["C4"]),
                sds,
                direction,
                lane,
                polarity,
                category,
                partition,
                die,
            )
            net_records.append(record)
            excel_records.append(
                ExcelRecord(
                    record,
                    " ".join(endpoint_pins["BGA"]),
                    " ".join(endpoint_pins["C4"]),
                )
            )
        if not net_records:
            raise NoMatchingSignals(
                "No SerDes SDS TX/RX lane networks were found in {0}".format(path)
            )
        validate_differential_pairs(net_records)
        return excel_records
    finally:
        edb.close()
        # PyEDB 2026.1 may leave this zero-byte lock marker even for a
        # read-only database.  It is runtime debris, not distributable input.
        (path / "edb.def.tmp").unlink(missing_ok=True)


def row_sort_key(record):
    net = record.net
    return (
        net.die,
        (0, net.sds) if net.sds is not None else (1, 0),
        DIRECTION_RANK[net.direction],
        net.lane,
        POLARITY_RANK[net.polarity],
        natural_key(net.name),
    )


def records_fingerprint(records):
    digest = hashlib.sha256()
    digest.update(FORMAT_VERSION.encode("ascii"))
    for record in sorted(records, key=row_sort_key):
        for value in record.values():
            digest.update(str(value).encode("utf-8"))
            digest.update(b"\0")
    return digest.hexdigest().upper()


def output_is_current(path, fingerprint):
    if not path.is_file():
        return False
    try:
        workbook = load_workbook(path, read_only=True)
        current = workbook.properties.keywords == "source-sha256=" + fingerprint
        workbook.close()
        return current
    except Exception:
        return False


def write_sheet(workbook, title, records):
    sheet = workbook.create_sheet(title)
    sheet.append(HEADERS)
    for record in records:
        sheet.append(record.values())
    style_netlist_sheet(sheet, records)


def create_workbook(records, fingerprint):
    workbook = Workbook()
    workbook.remove(workbook.active)
    ordered = sorted(records, key=row_sort_key)
    write_sheet(workbook, "ALL", ordered)
    write_sheet(workbook, "RX", [row for row in ordered if row.net.direction == "RX"])
    write_sheet(workbook, "TX", [row for row in ordered if row.net.direction == "TX"])
    workbook.properties.title = "Ketupa SerDes Differential Netlist Classification"
    workbook.properties.keywords = "source-sha256=" + fingerprint
    return workbook


def verify_workbook(path, records):
    workbook = load_workbook(path, read_only=False, data_only=True)
    expected_sheets = ["ALL", "RX", "TX"]
    if workbook.sheetnames != expected_sheets:
        raise AssertionError(
            "Unexpected worksheet order: {0}".format(workbook.sheetnames)
        )
    all_rows = list(workbook["ALL"].iter_rows(values_only=True))
    if all_rows[0] != HEADERS:
        raise AssertionError("The six-column workbook contract changed")
    expected = [row.values() for row in sorted(records, key=row_sort_key)]
    if all_rows[1:] != expected:
        raise AssertionError("ALL worksheet rows do not match the parsed report")
    for direction_sheet in workbook.sheetnames[1:]:
        actual = list(workbook[direction_sheet].iter_rows(min_row=2, values_only=True))
        wanted = [
            row.values()
            for row in sorted(records, key=row_sort_key)
            if row.net.direction == direction_sheet
        ]
        if actual != wanted:
            raise AssertionError(
                "Worksheet {0} is not grouped correctly".format(direction_sheet)
            )
    ordered = sorted(records, key=row_sort_key)
    for sheet_name in workbook.sheetnames:
        sheet_records = (
            ordered
            if sheet_name == "ALL"
            else [row for row in ordered if row.net.direction == sheet_name]
        )
        blocks = []
        for record in sheet_records:
            key = color_block_key(record)
            if not blocks or blocks[-1][0] != key:
                blocks.append((key, []))
            blocks[-1][1].append(record)
        row_number = 2
        previous_color = None
        for key, block_records in blocks:
            start_row = row_number
            end_row = start_row + len(block_records) - 1
            colors = {
                workbook[sheet_name].cell(row, column).fill.fgColor.rgb
                for row in range(start_row, end_row + 1)
                for column in range(1, 6)
            }
            if len(colors) != 1:
                raise AssertionError(
                    "Worksheet {0} block {1} does not use one background color".format(
                        sheet_name, key
                    )
                )
            color = next(iter(colors))
            if color == previous_color:
                raise AssertionError(
                    "Worksheet {0} adjacent blocks share one color: {1}".format(
                        sheet_name, key
                    )
                )
            if workbook[sheet_name].cell(start_row, 1).border.top.style != "medium":
                raise AssertionError(
                    "Worksheet {0} block {1} has no top separator".format(
                        sheet_name, key
                    )
                )
            if workbook[sheet_name].cell(end_row, 1).border.bottom.style != "medium":
                raise AssertionError(
                    "Worksheet {0} block {1} has no bottom separator".format(
                        sheet_name, key
                    )
                )
            previous_color = color
            row_number = end_row + 1
    workbook.close()


def _publish_excel_records(records, output_path, force=False):
    output_path = Path(output_path).resolve()
    fingerprint = records_fingerprint(records)
    if not force and output_is_current(output_path, fingerprint):
        print("Unchanged: {0}".format(output_path))
        return output_path
    temporary = output_path.with_name(".{0}.tmp.xlsx".format(output_path.stem))
    workbook = create_workbook(records, fingerprint)
    try:
        workbook.save(temporary)
        verify_workbook(temporary, records)
        temporary.replace(output_path)
    finally:
        workbook.close()
        if temporary.exists():
            temporary.unlink()
    counts = Counter(record.net.direction for record in records)
    print("Created: {0}".format(output_path))
    print(
        "Networks: {0}; sheets: {1}".format(
            len(records),
            "ALL, RX, TX",
        )
    )
    print(
        "Directions: "
        + ", ".join(
            "{0}={1}".format(direction, counts[direction])
            for direction in ("RX", "TX")
        )
    )
    print("Output order: {0}".format(SORT_CONTRACT))
    return output_path


def _legacy_convert(input_path, output_path=None, groups="auto", force=False):
    input_path = Path(input_path).resolve()
    if (
        input_path.suffix.casefold() not in {".htm", ".html"}
        or not input_path.is_file()
    ):
        raise ValueError(
            "Input must be an existing .htm or .html file: {0}".format(input_path)
        )
    output_path = canonical_xlsx_path(input_path, output_path)
    records = build_excel_records(
        load_net_records(input_path, parse_group_selection(groups))
    )
    return _publish_excel_records(records, output_path, force)


def convert_aedb(input_path, output_path=None, groups="auto", force=False):
    """Create the normal six-column workbook from a packaged native AEDB."""
    input_path = Path(input_path).resolve()
    output_path = canonical_xlsx_path(input_path, output_path)
    records = load_aedb_excel_records(
        input_path, parse_group_selection(groups)
    )
    return _publish_excel_records(records, output_path, force)


def convert(input_path, output_path=None, groups="auto", force=False):
    """Keep the qualified legacy path and use the approved adapter on rejection."""
    input_path = Path(input_path).resolve()
    selected_output = canonical_xlsx_path(input_path, output_path)
    from script.data_adapter.runtime import convert_with_fallback

    return convert_with_fallback(
        lambda: _legacy_convert(input_path, selected_output, groups, force),
        input_path,
        selected_output,
        family="serdes",
        topology="pkg",
        groups=groups,
        force=force,
    )


def convert_source(
    input_path,
    output_path=None,
    groups="auto",
    force=False,
):
    """Convert one v1/v2 report or generate a temporary v2 from MCM/SIP."""
    input_path = Path(input_path).resolve()
    if input_path.suffix.casefold() == ".aedb":
        output_path = Path(
            output_path
            or default_design_output(PROJECT_ROOT, input_path, "netlist")
        ).resolve()
        return convert_aedb(input_path, output_path, groups, force)
    if input_path.suffix.casefold() in {".mcm", ".sip"}:
        output_path = Path(
            output_path
            or default_design_output(PROJECT_ROOT, input_path, "netlist")
        ).resolve()
        with temporary_v2_report(
            input_path,
            "netlist",
            PROJECT_ROOT,
        ) as report_path:
            result = convert(report_path, output_path, groups, force)
        print("Removed temporary v2 report.")
        return result
    return convert(input_path, output_path, groups, force)


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "input",
        nargs="?",
        type=Path,
        default=PROJECT_ROOT,
        help="one _v1/_v2 report, AEDB, MCM, or SIP; default: active PKG folder",
    )
    parser.add_argument("-o", "--output", type=Path)
    parser.add_argument(
        "--groups",
        "--lanes",
        dest="groups",
        default="auto",
        help="SerDes lanes: auto, 0,2, or 0-15",
    )
    parser.add_argument("--force", action="store_true")
    return parser.parse_args()


def main():
    args = parse_args()
    input_path = args.input.resolve()
    if input_path.suffix.casefold() in {
        ".htm",
        ".html",
        ".aedb",
        ".mcm",
        ".sip",
    }:
        try:
            convert_source(
                input_path,
                args.output,
                args.groups,
                args.force,
            )
        except NoMatchingSignals as error:
            print("[SKIPPED] {0}: {1}".format(input_path.name, error))
            return NO_SIGNALS
        except Exception as error:
            print("[FAILED] {0}: {1}".format(input_path.name, error))
            return 1
        return 0

    input_path.mkdir(parents=True, exist_ok=True)
    reports = sorted(
        (
            path
            for path in input_path.iterdir()
            if path.is_file() and path.suffix.casefold() in {".htm", ".html"}
        ),
        key=lambda path: path.name.casefold(),
    )
    if not reports:
        print("No netlist .htm/.html files found in: {0}".format(input_path))
        print("Put _v1.htm or _v2.htm files there, then run script/01_Netlist.sh again.")
        return 0

    failures = 0
    skipped = 0
    for report in reports:
        print("[NETLIST] {0}".format(report.name))
        try:
            convert_source(report, groups=args.groups, force=True)
        except NoMatchingSignals as error:
            skipped += 1
            print("[SKIPPED] {0}: {1}".format(report.name, error))
        except Exception as error:
            failures += 1
            print("[FAILED] {0}: {1}".format(report.name, error))
    print(
        "Netlist reports: {0}  Succeeded: {1}  Failed: {2}".format(
            len(reports), len(reports) - failures - skipped, failures
        )
    )
    print("Skipped: {0}".format(skipped))
    return 1 if failures else NO_SIGNALS if skipped == len(reports) else 0


if __name__ == "__main__":
    raise SystemExit(extractor_exit(main))
