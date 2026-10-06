#!/usr/bin/env python3
"""Convert a DDR netlist v1/v2 report or BRD/MCM design into Excel."""

from __future__ import annotations

import argparse
import csv
import hashlib
import os
import re
import subprocess
import tempfile
from collections import Counter
from contextlib import contextmanager
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from artifact_paths import artifact_path, NO_SIGNALS, NoMatchingSignals, extractor_exit
from cadence_runtime import resolve_cadence_export_command as _resolve_cadence_export_command, run_cadence_export

FORMAT_VERSION = "ketupa-netlist-v10"
PROJECT_ROOT = Path(os.environ["KETUPA_WORKFLOW_INPUT"]).resolve()
DESIGN_SUFFIXES = {".brd", ".mcm"}
REPORT_CODES = {"netlist": "net", "component": "cmp"}
REPORT_PREFIXES = {
    "netlist": "NetListReport",
    "component": "ComponentReport",
}
VARIANT_SUFFIX = re.compile(r"_v(?P<version>[12])$", re.I)
SIGNAL_ORDER = (
    "DQ",
    "DM",
    "DQS",
    "WCK",
    "CA",
    "CS",
    "CKE",
    "CK_C",
    "CK_T",
    "CK",
    "ODT",
    "RESET",
    "ALERT",
    "TEN",
    "ZQ",
)
SIGNAL_RANK = {name: index for index, name in enumerate(SIGNAL_ORDER)}
SEPARATOR = r"[_\-.]+"
NET_PATTERNS = (
    re.compile(
        rf"^M(?P<group>\d+){SEPARATOR}CH(?:ANNEL)?(?P<channel>[A-Z0-9]+){SEPARATOR}(?P<signal>.+)$",
        re.I,
    ),
    re.compile(
        rf"^CH(?:ANNEL)?(?P<channel>[A-Z0-9]+){SEPARATOR}M(?P<group>\d+){SEPARATOR}(?P<signal>.+)$",
        re.I,
    ),
    re.compile(
        rf"^M(?P<group>\d+){SEPARATOR}(?P<signal>.+){SEPARATOR}CH(?:ANNEL)?(?P<channel>[A-Z0-9]+)$",
        re.I,
    ),
    re.compile(
        rf"^M(?P<group>\d+){SEPARATOR}(?P<signal>.+){SEPARATOR}(?P<channel>[A-D])$",
        re.I,
    ),
)
CHIP_SUFFIX = re.compile(r"(?:_|-|\.)G(?P<index>\d*)$", re.I)
PIN = re.compile(r"^(?P<refdes>[^.\s]+)\.(?P<pin>\S+)$")
HEADERS = ("Net Name", "PIN 1", "PIN 2", "Signal Type", "M Group", "Chip Partition")

HEADER_COLOR = "17365D"
HEADER_ACCENT = "4472C4"
TEXT_COLOR = "243447"
GRID_COLOR = "D6DEE8"
BLOCK_COLORS = (
    "D9EAF7",
    "E2F0D9",
    "FCE4D6",
    "E4DFEC",
    "FFF2CC",
    "DDEBF7",
    "EADCF8",
    "F4CCCC",
    "D9EAD3",
    "D0E0E3",
    "FCE5CD",
    "DDEBF1",
)
CHIP_COLORS = (
    "E2F0D9",
    "FCE4D6",
    "DDEBF7",
    "E4DFEC",
    "FFF2CC",
    "D9EAD3",
)
CHANNEL_ACCENTS = ("4472C4", "70AD47", "ED7D31", "A64D79", "5B9BD5")


class TableParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.row = None
        self.cell = None
        self.rows = []

    def handle_starttag(self, tag, attrs):
        del attrs
        tag = tag.casefold()
        if tag == "tr":
            self.row = []
        elif tag in {"td", "th"} and self.row is not None:
            self.cell = []

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)

    def handle_endtag(self, tag):
        tag = tag.casefold()
        if tag in {"td", "th"} and self.cell is not None:
            self.row.append(" ".join("".join(self.cell).replace("\xa0", " ").split()))
            self.cell = None
        elif tag == "tr" and self.row is not None:
            if self.row:
                self.rows.append(self.row)
            self.row = None


def _clean_cell(value):
    return " ".join(str(value).replace("\xa0", " ").split())


def _decode_report(path):
    raw = Path(path).read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("Cannot decode Cadence report: {0}".format(path))


def report_variant(path, text=None):
    path = Path(path)
    text = _decode_report(path) if text is None else text
    content_variant = "v1" if "<table" in text.casefold() else "v2"
    suffix_match = VARIANT_SUFFIX.search(path.stem)
    if suffix_match is not None:
        named_variant = "v" + suffix_match.group("version")
        if named_variant != content_variant:
            raise ValueError(
                "Report suffix/content mismatch for {0}: name={1}, content={2}".format(
                    path, named_variant, content_variant
                )
            )
    return content_variant


def canonical_report_stem(path):
    return VARIANT_SUFFIX.sub("", Path(path).stem)


def canonical_xlsx_path(path, output_path=None):
    path = Path(path).resolve()
    selected = (
        Path(output_path).resolve()
        if output_path is not None
        else path.with_name(canonical_report_stem(path) + ".xlsx")
    )
    if selected.suffix.casefold() != ".xlsx":
        raise ValueError("Output must use .xlsx: {0}".format(selected))
    return selected


def _records_from_rows(rows, required_headers):
    requested = tuple(str(header).strip() for header in required_headers)
    required = {header.casefold() for header in requested}
    for index, row in enumerate(rows):
        normalized = [_clean_cell(value).casefold() for value in row]
        if not required.issubset(set(normalized)):
            continue
        positions = {
            header: normalized.index(header.casefold()) for header in requested
        }
        last_column = max(positions.values())
        records = []
        for values in rows[index + 1 :]:
            cleaned = [_clean_cell(value) for value in values]
            if len(cleaned) <= last_column or not any(cleaned):
                continue
            records.append(
                {header: cleaned[column] for header, column in positions.items()}
            )
        if records:
            return records
    raise ValueError(
        "Cadence report does not contain the required columns: {0}".format(
            ", ".join(requested)
        )
    )


def read_report_table(path, required_headers):
    path = Path(path).resolve()
    if path.suffix.casefold() not in {".htm", ".html"} or not path.is_file():
        raise ValueError("Report must be an existing .htm/.html file: {0}".format(path))
    text = _decode_report(path)
    if report_variant(path, text) == "v1":
        parser = TableParser()
        parser.feed(text)
        rows = parser.rows
    else:
        rows = list(csv.reader(text.splitlines()))
    return _records_from_rows(rows, required_headers)


def resolve_cadence_export_command():
    return _resolve_cadence_export_command()


def default_design_output(project_root, design_path, report_kind):
    return artifact_path(project_root, design_path, report_kind, "lpddr")


@contextmanager
def temporary_v2_report(design_path, report_kind, project_root=None):
    del project_root
    design_path = Path(design_path).resolve()
    if design_path.suffix.casefold() not in DESIGN_SUFFIXES or not design_path.is_file():
        raise ValueError("Design must be an existing .brd/.mcm file: {0}".format(design_path))
    with tempfile.TemporaryDirectory(prefix="ketupa_cadence_report_") as directory:
        report_path = Path(directory) / "{0}_{1}_v2.htm".format(
            REPORT_PREFIXES[report_kind], design_path.stem
        )
        run_cadence_export(REPORT_CODES[report_kind], design_path, report_path)
        report_variant(report_path)
        try:
            yield report_path
        finally:
            pass


def read_table(path, required_headers):
    return read_report_table(path, required_headers)


def _rgb(value):
    value = value.lstrip("#")
    return tuple(int(value[index : index + 2], 16) for index in (0, 2, 4))


def _hex(values):
    return "".join("{0:02X}".format(max(0, min(255, value))) for value in values)


def _darken(color, amount=0.28):
    return _hex(tuple(round(value * (1 - amount)) for value in _rgb(color)))


def _channel_color(channel):
    number = sum(ord(character) for character in str(channel).upper())
    return CHANNEL_ACCENTS[number % len(CHANNEL_ACCENTS)]


def color_block_key(record):
    net = record.net
    return net.chip, net.group, net.channel


def style_netlist_sheet(sheet, records):
    for cell in sheet[1][:6]:
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_COLOR)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=Side(style="medium", color=HEADER_ACCENT))
    sheet.row_dimensions[1].height = 28
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:F{0}".format(sheet.max_row)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90
    for column, width in zip("ABCDEF", (34, 30, 44, 14, 12, 18)):
        sheet.column_dimensions[column].width = width

    block_keys = list(dict.fromkeys(color_block_key(record) for record in records))
    block_colors = {
        key: BLOCK_COLORS[index % len(BLOCK_COLORS)]
        for index, key in enumerate(block_keys)
    }
    previous = None
    for row_index, record in enumerate(records, start=2):
        net = record.net
        block_color = block_colors[color_block_key(record)]
        chip_color = CHIP_COLORS[(net.chip - 1) % len(CHIP_COLORS)]
        channel_color = _channel_color(net.channel)
        chip_changed = previous is None or previous.net.chip != net.chip
        group_changed = chip_changed or previous.net.group != net.group
        block_changed = group_changed or previous.net.channel != net.channel
        next_record = records[row_index - 1] if row_index - 1 < len(records) else None
        block_ends = next_record is None or color_block_key(
            next_record
        ) != color_block_key(record)
        top_color = _darken(chip_color) if chip_changed else _darken(block_color)
        for column_index in range(1, 7):
            cell = sheet.cell(row_index, column_index)
            cell.font = Font(
                name="Microsoft YaHei",
                size=9,
                color=TEXT_COLOR,
                bold=column_index in {1, 5, 6},
            )
            cell.fill = PatternFill(
                "solid", fgColor=chip_color if column_index == 6 else block_color
            )
            cell.alignment = Alignment(
                horizontal="center" if column_index >= 4 else "left",
                vertical="center",
            )
            cell.border = Border(
                left=(
                    Side(style="medium", color=channel_color)
                    if column_index == 1
                    else Side()
                ),
                top=(
                    Side(
                        style="medium",
                        color=top_color
                        if chip_changed or group_changed
                        else channel_color,
                    )
                    if block_changed
                    else Side()
                ),
                bottom=Side(
                    style="medium" if block_ends else "hair",
                    color=_darken(block_color) if block_ends else GRID_COLOR,
                ),
            )
        sheet.row_dimensions[row_index].height = 19
        previous = record
    sheet.sheet_properties.tabColor = (
        HEADER_COLOR
        if sheet.title == "All"
        else _darken(BLOCK_COLORS[records[0].net.group % len(BLOCK_COLORS)], 0.15)
    )


@dataclass(frozen=True)
class NetRecord:
    name: str
    pins: tuple[str, ...]
    group: int
    channel: str
    signal: str
    signal_type: str
    partition: str
    chip: int


@dataclass(frozen=True)
class ExcelRecord:
    net: NetRecord
    pin1: str
    pin2: str

    def values(self):
        return (
            self.net.name,
            self.pin1,
            self.pin2,
            self.net.signal_type,
            "M{0}".format(self.net.group),
            self.net.partition,
        )


def natural_key(value):
    return tuple(
        (0, int(part)) if part.isdigit() else (1, part.casefold())
        for part in re.split(r"(\d+)", str(value))
        if part
    )


def refdes(pin):
    match = PIN.fullmatch(pin)
    if not match:
        raise ValueError("Invalid RefDes.Pin token: {0}".format(pin))
    return match.group("refdes")


def signal_type(signal):
    upper = signal.upper()
    tokens = [token for token in re.split(r"[^A-Z0-9]+", upper) if token]

    def contains(*prefixes):
        return any(token.startswith(prefixes) for token in tokens)

    if contains("RDQS", "DQS"):
        return "DQS"
    if contains("WCK"):
        return "WCK"
    if contains("CKE"):
        return "CKE"
    if contains("CS"):
        return "CS"
    if contains("CK"):
        token_set = set(tokens)
        if token_set.intersection({"C", "N", "DN"}) or upper.endswith("_C"):
            return "CK_C"
        if token_set.intersection({"T", "P", "DP"}) or upper.endswith("_T"):
            return "CK_T"
        return "CK"
    if contains("DMI", "DM"):
        return "DM"
    if contains("DQ"):
        return "DQ"
    if contains("CA"):
        return "CA"
    if contains("ODT"):
        return "ODT"
    if contains("RESET", "RST"):
        return "RESET"
    if contains("ALERT"):
        return "ALERT"
    if contains("TEN"):
        return "TEN"
    if contains("ZQ"):
        return "ZQ"
    return None


def parse_net_name(name):
    clean = name.strip()
    suffix = CHIP_SUFFIX.search(clean)
    chip = 1
    classified = clean
    if suffix:
        suffix_index = int(suffix.group("index") or "1")
        chip = suffix_index + 1
        classified = clean[: suffix.start()]
    for pattern in NET_PATTERNS:
        match = pattern.fullmatch(classified)
        if match:
            signal = match.group("signal").upper()
            category = signal_type(signal)
            if category is None:
                return None
            return (
                int(match.group("group")),
                match.group("channel").upper(),
                signal,
                category,
                "CHIP_{0}".format(chip),
                chip,
            )
    return None


def chip_number(value):
    match = CHIP_SUFFIX.search(str(value).strip())
    return 1 if match is None else int(match.group("index") or "1") + 1


def parse_group_selection(value):
    text = str(value).strip().casefold()
    if text in {"", "auto", "all"}:
        return None
    selected = set()
    for token in text.split(","):
        token = token.strip().removeprefix("m")
        if "-" in token:
            start, end = token.split("-", 1)
            start_number = int(start.removeprefix("m"))
            end_number = int(end.removeprefix("m"))
            if start_number > end_number:
                raise ValueError("Descending M-group range: {0}".format(token))
            selected.update(range(start_number, end_number + 1))
        else:
            selected.add(int(token))
    return selected


def load_net_records(path, selected_groups=None):
    records = []
    seen = set()
    for row in read_table(path, ("Net Name", "Net Pins")):
        name = row["Net Name"].strip()
        parsed = parse_net_name(name)
        if parsed is None:
            continue
        group, channel, signal, category, partition, chip = parsed
        if selected_groups is not None and group not in selected_groups:
            continue
        key = name.casefold()
        if key in seen:
            raise ValueError("Duplicate DDR network: {0}".format(name))
        seen.add(key)
        pins = tuple(
            token for token in re.split(r"[\s,;|]+", row["Net Pins"].strip()) if token
        )
        if len(pins) < 2:
            raise ValueError("DDR network requires at least two pins: {0}".format(name))
        mismatched = []
        for pin in pins:
            component = refdes(pin)
            if chip_number(component) != chip:
                mismatched.append(component)
        mismatched = sorted(set(mismatched), key=natural_key)
        if mismatched:
            raise ValueError(
                "Network/RefDes chip partition mismatch for {0}: {1}".format(
                    name, ", ".join(mismatched)
                )
            )
        records.append(
            NetRecord(name, pins, group, channel, signal, category, partition, chip)
        )
    if not records:
        raise NoMatchingSignals("No M*/CHA-CHD DDR networks were found in {0}".format(path))
    return records


def primary_refdes_by_partition(records):
    partitions = sorted({record.partition for record in records}, key=natural_key)
    coverage = {partition: Counter() for partition in partitions}
    first_seen = {partition: {} for partition in partitions}
    for row_number, record in enumerate(records):
        for pin_index, component in enumerate(
            dict.fromkeys(refdes(pin) for pin in record.pins)
        ):
            coverage[record.partition][component] += 1
            first_seen[record.partition].setdefault(
                component,
                (row_number, pin_index),
            )
    return {
        partition: min(
            counts,
            key=lambda name: (
                -counts[name],
                first_seen[partition][name],
                natural_key(name),
            ),
        )
        for partition, counts in coverage.items()
        if counts
    }


def build_excel_records(records):
    primary = primary_refdes_by_partition(records)
    output = []
    for record in records:
        controller = primary[record.partition]
        controller_pins = [pin for pin in record.pins if refdes(pin) == controller]
        target_pins = [pin for pin in record.pins if refdes(pin) != controller]
        if not controller_pins or not target_pins:
            raise ValueError(
                "Cannot split controller and DDR endpoints for {0}".format(record.name)
            )
        output.append(
            ExcelRecord(record, " ".join(controller_pins), " ".join(target_pins))
        )
    return output


def row_sort_key(record):
    net = record.net
    return (
        net.chip,
        net.group,
        natural_key(net.channel),
        SIGNAL_RANK.get(net.signal_type, len(SIGNAL_RANK)),
        natural_key(net.signal),
        natural_key(net.name),
    )


def records_fingerprint(records):
    digest = hashlib.sha256()
    digest.update(FORMAT_VERSION.encode("ascii"))
    for record in sorted(records, key=row_sort_key):
        for value in record.values():
            digest.update(str(value).encode("utf-8"))
            digest.update(b"\0")
    return digest.hexdigest().upper()


def output_is_current(path, fingerprint):
    if not path.is_file():
        return False
    try:
        workbook = load_workbook(path, read_only=True)
        current = workbook.properties.keywords == "source-sha256=" + fingerprint
        workbook.close()
        return current
    except Exception:
        return False


def write_sheet(workbook, title, records):
    sheet = workbook.create_sheet(title)
    sheet.append(HEADERS)
    for record in records:
        sheet.append(record.values())
    style_netlist_sheet(sheet, records)


def create_workbook(records, fingerprint):
    workbook = Workbook()
    workbook.remove(workbook.active)
    ordered = sorted(records, key=row_sort_key)
    write_sheet(workbook, "All", ordered)
    for group in sorted({record.net.group for record in ordered}):
        write_sheet(
            workbook,
            "M{0}".format(group),
            [record for record in ordered if record.net.group == group],
        )
    workbook.properties.title = "Ketupa DDR Netlist Classification"
    workbook.properties.keywords = "source-sha256=" + fingerprint
    return workbook


def verify_workbook(path, records):
    workbook = load_workbook(path, read_only=False, data_only=True)
    expected_sheets = ["All"] + [
        "M{0}".format(group) for group in sorted({row.net.group for row in records})
    ]
    if workbook.sheetnames != expected_sheets:
        raise AssertionError(
            "Unexpected worksheet order: {0}".format(workbook.sheetnames)
        )
    all_rows = list(workbook["All"].iter_rows(values_only=True))
    if all_rows[0] != HEADERS:
        raise AssertionError("The six-column workbook contract changed")
    expected = [row.values() for row in sorted(records, key=row_sort_key)]
    if all_rows[1:] != expected:
        raise AssertionError("All worksheet rows do not match the parsed HTML")
    for group_sheet in workbook.sheetnames[1:]:
        group_number = int(group_sheet[1:])
        actual = list(workbook[group_sheet].iter_rows(min_row=2, values_only=True))
        wanted = [
            row.values()
            for row in sorted(records, key=row_sort_key)
            if row.net.group == group_number
        ]
        if actual != wanted:
            raise AssertionError(
                "Worksheet {0} is not grouped correctly".format(group_sheet)
            )
    ordered = sorted(records, key=row_sort_key)
    for sheet_name in workbook.sheetnames:
        sheet_records = (
            ordered
            if sheet_name == "All"
            else [row for row in ordered if row.net.group == int(sheet_name[1:])]
        )
        blocks = []
        for record in sheet_records:
            key = color_block_key(record)
            if not blocks or blocks[-1][0] != key:
                blocks.append((key, []))
            blocks[-1][1].append(record)
        row_number = 2
        previous_color = None
        for key, block_records in blocks:
            start_row = row_number
            end_row = start_row + len(block_records) - 1
            colors = {
                workbook[sheet_name].cell(row, column).fill.fgColor.rgb
                for row in range(start_row, end_row + 1)
                for column in range(1, 6)
            }
            if len(colors) != 1:
                raise AssertionError(
                    "Worksheet {0} block {1} does not use one background color".format(
                        sheet_name, key
                    )
                )
            color = next(iter(colors))
            if color == previous_color:
                raise AssertionError(
                    "Worksheet {0} adjacent blocks share one color: {1}".format(
                        sheet_name, key
                    )
                )
            if workbook[sheet_name].cell(start_row, 1).border.top.style != "medium":
                raise AssertionError(
                    "Worksheet {0} block {1} has no top separator".format(
                        sheet_name, key
                    )
                )
            if workbook[sheet_name].cell(end_row, 1).border.bottom.style != "medium":
                raise AssertionError(
                    "Worksheet {0} block {1} has no bottom separator".format(
                        sheet_name, key
                    )
                )
            previous_color = color
            row_number = end_row + 1
    workbook.close()


def _legacy_convert(input_path, output_path=None, groups="auto", force=False):
    input_path = Path(input_path).resolve()
    if (
        input_path.suffix.casefold() not in {".htm", ".html"}
        or not input_path.is_file()
    ):
        raise ValueError(
            "Input must be an existing .htm or .html file: {0}".format(input_path)
        )
    output_path = canonical_xlsx_path(input_path, output_path)
    records = build_excel_records(
        load_net_records(input_path, parse_group_selection(groups))
    )
    fingerprint = records_fingerprint(records)
    if not force and output_is_current(output_path, fingerprint):
        print("Unchanged: {0}".format(output_path))
        return output_path
    temporary = output_path.with_name(".{0}.tmp.xlsx".format(output_path.stem))
    workbook = create_workbook(records, fingerprint)
    try:
        workbook.save(temporary)
        verify_workbook(temporary, records)
        temporary.replace(output_path)
    finally:
        workbook.close()
        if temporary.exists():
            temporary.unlink()
    counts = Counter(record.net.channel for record in records)
    print("Created: {0}".format(output_path))
    print(
        "Networks: {0}; sheets: {1}".format(
            len(records),
            ", ".join(
                ["All"]
                + [
                    "M{0}".format(group)
                    for group in sorted({row.net.group for row in records})
                ]
            ),
        )
    )
    print(
        "Channels: "
        + ", ".join(
            "CH{0}={1}".format(channel, counts[channel])
            for channel in sorted(counts, key=natural_key)
        )
    )
    return output_path


def convert(input_path, output_path=None, groups="auto", force=False):
    """Keep the qualified legacy path and use the approved adapter on rejection."""
    input_path = Path(input_path).resolve()
    selected_output = canonical_xlsx_path(input_path, output_path)
    from script.data_adapter.runtime import convert_with_fallback

    return convert_with_fallback(
        lambda: _legacy_convert(input_path, selected_output, groups, force),
        input_path,
        selected_output,
        family="lpddr",
        topology="pcb",
        groups=groups,
        force=force,
    )


def convert_source(
    input_path,
    output_path=None,
    groups="auto",
    force=False,
):
    """Convert one v1/v2 report or generate a temporary v2 from BRD/MCM."""
    input_path = Path(input_path).resolve()
    if input_path.suffix.casefold() in DESIGN_SUFFIXES:
        output_path = Path(
            output_path
            or default_design_output(PROJECT_ROOT, input_path, "netlist")
        ).resolve()
        with temporary_v2_report(
            input_path,
            "netlist",
            PROJECT_ROOT,
        ) as report_path:
            result = convert(report_path, output_path, groups, force)
        print("Removed temporary v2 report.")
        return result
    return convert(input_path, output_path, groups, force)


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "input",
        nargs="?",
        type=Path,
        default=PROJECT_ROOT,
        help="one _v1/_v2 report or BRD; default: active PCB folder",
    )
    parser.add_argument("-o", "--output", type=Path)
    parser.add_argument("--groups", default="auto", help="auto, M0,M2, or M0-M99")
    parser.add_argument("--force", action="store_true")
    return parser.parse_args()


def main():
    args = parse_args()
    input_path = args.input.resolve()
    if input_path.suffix.casefold() in {".htm", ".html", ".brd", ".mcm"}:
        convert_source(
            input_path,
            args.output,
            args.groups,
            args.force,
        )
        return 0

    input_path.mkdir(parents=True, exist_ok=True)
    reports = sorted(
        (
            path
            for path in input_path.iterdir()
            if path.is_file() and path.suffix.casefold() in {".htm", ".html"}
        ),
        key=lambda path: path.name.casefold(),
    )
    if not reports:
        print("No netlist .htm/.html files found in: {0}".format(input_path))
        print("Put _v1.htm or _v2.htm files there, then run script/01_Netlist.sh again.")
        return 0

    failures = 0
    skipped = 0
    for report in reports:
        print("[NETLIST] {0}".format(report.name))
        try:
            convert_source(report, groups=args.groups, force=True)
        except NoMatchingSignals as error:
            skipped += 1
            print("[SKIPPED] {0}: {1}".format(report.name, error))
        except Exception as error:
            failures += 1
            print("[FAILED] {0}: {1}".format(report.name, error))
    print(
        "Netlist reports: {0}  Succeeded: {1}  Failed: {2}".format(
            len(reports), len(reports) - failures - skipped, failures
        )
    )
    print("Skipped: {0}".format(skipped))
    return 1 if failures else NO_SIGNALS if skipped == len(reports) else 0


if __name__ == "__main__":
    raise SystemExit(extractor_exit(main))
